from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from create_legal_tz_docx import (
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    MUTED,
    PALE_YELLOW,
    add_callout,
    add_hyperlink,
    add_list_item,
    add_meta_line,
    add_numbering_definitions,
    add_page_number,
    add_table,
    configure_styles,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Краткое_ТЗ_юристу_RFID_панель_Hello_Park.docx"


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.right_margin = Inches(0.9)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    configure_styles(doc)

    # Slightly tighter rhythm for a short working brief.
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.16
    for style_name in ("Heading 1", "Heading 2"):
        doc.styles[style_name].paragraph_format.space_before = Pt(10)
        doc.styles[style_name].paragraph_format.space_after = Pt(5)

    bullet_num, decimal_num = add_numbering_definitions(doc)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("HELLO PARK  |  КРАТКОЕ ТЗ ЮРИСТУ ПО RFID-ПАНЕЛИ")
    set_run_font(run, size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("КРАТКОЕ ТЕХНИЧЕСКОЕ ЗАДАНИЕ ЮРИСТУ")
    set_run_font(run, size=10.5, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Подготовка парков к внедрению RFID-панели")
    set_run_font(run, size=23, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Регистрация гостей, данные детей, RFID-браслеты и программа лояльности Hello Park")
    set_run_font(run, size=11.5, color=MUTED)

    add_meta_line(doc, "Цель", "Проверить юридическую готовность каждого парка и подготовить простой комплект документов для запуска")
    add_meta_line(doc, "Версия", "Краткий рабочий черновик от 18.08.2026")

    add_callout(
        doc,
        "Что нужно получить.",
        "Понятный ответ по каждому парку: кто работает с данными, какие документы и согласия нужны, требуется ли уведомление регулятора и что необходимо изменить в форме регистрации до запуска.",
    )

    doc.add_heading("1. Что мы собираемся внедрить", level=1)
    doc.add_paragraph(
        "Во всех парках планируется использовать единую веб-панель для регистрации семей и привязки RFID-браслетов к профилям детей. Браслет нужен, чтобы ребенок мог играть на интерактивных аттракционах, сохранять игровой прогресс, получать баллы и призы, а при следующем визите продолжить игру с прежнего профиля."
    )

    doc.add_heading("2. Пользовательский путь", level=1)
    steps = [
        "Гость сканирует QR-код в парке и открывает мобильную форму регистрации.",
        "Вводит номер телефона и подтверждает его по SMS.",
        "Знакомится с документами своего парка и подтверждает необходимые согласия.",
        "Указывает ФИО взрослого и добавляет детей: имя и дату рождения.",
        "После оплаты билета кассир находит профиль и привязывает RFID-браслет к выбранному ребенку.",
        "Во время посещения сохраняются игровой прогресс, баллы, призы и история визитов. При повторном посещении профиль можно найти и привязать новый браслет.",
    ]
    for step in steps:
        add_list_item(doc, step, decimal_num)

    doc.add_heading("3. Что именно мы собираем и храним", level=1)
    data_rows = [
        ("Взрослый", "ФИО, номер телефона, подтверждение телефона"),
        ("Ребенок", "Имя, дата рождения, связь с профилем взрослого"),
        ("RFID", "UID браслета, к какому ребенку привязан, замена и очистка"),
        ("Посещения", "Парк, дата и время визита"),
        ("Игровой профиль", "Аватар, прогресс, достижения, уровень"),
        ("Лояльность", "Баллы, доступные и выданные призы"),
        ("Подтверждения", "Какие документы и версии показаны гостю, что он подтвердил, дата и время"),
    ]
    data_table = add_table(doc, ["Блок", "Данные"], data_rows, [2500, 6860], font_size=9.4)
    tr_pr = data_table.rows[0]._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)

    add_callout(
        doc,
        "Отдельно проверить.",
        "Данные ребенка вводит взрослый. Нужна понятная формулировка, что он является родителем или законным представителем и имеет право предоставить эти данные.",
        fill=PALE_YELLOW,
    )

    doc.add_heading("4. Что мы хотим сделать", level=1)
    plans = [
        "Запустить одну общую систему, но показывать гостю документы и реквизиты именно того парка, в котором он находится.",
        "Хранить профиль семьи и игровой прогресс, чтобы ими можно было пользоваться при повторных посещениях.",
        "Ограничить доступ сотрудников данными своего парка, подключить нужные сервисы и сохранять доказательство согласий гостя.",
    ]
    for item in plans:
        add_list_item(doc, item, bullet_num)

    doc.add_heading("5. Задачи юриста", level=1)
    lawyer_tasks = [
        "Проверить правила посещения, политику обработки данных, согласия и программу лояльности по каждому парку на странице hello-park.ru/r.",
        "Определить по каждому парку юридическое лицо и его роль: является ли оно оператором персональных данных либо работает с данными по поручению центральной компании.",
        "Проверить, нужно ли каждому российскому парку подавать новое уведомление в Роскомнадзор или обновлять уже поданные сведения в связи с RFID-системой.",
        "Проверить данные детей и подготовить правильную формулировку для родителя или законного представителя.",
        "Подготовить тексты для формы регистрации: отдельное согласие на обработку данных, принятие правил парка и программы лояльности; отдельно - согласие на рекламу, если она будет.",
        "Определить сроки хранения, доступ, исправление и удаление данных; проверить договоры с платформой и подрядчиками. Для иностранных парков определить необходимость местного юриста.",
    ]
    _, lawyer_num = add_numbering_definitions(doc)
    for item in lawyer_tasks:
        add_list_item(doc, item, lawyer_num)

    p = doc.add_paragraph()
    p.add_run("Документы всех парков для проверки: ")
    add_hyperlink(p, "hello-park.ru/r", "https://hello-park.ru/r")

    heading = doc.add_heading("6. Что важно изменить в текущей форме", level=1)
    heading.paragraph_format.page_break_before = True
    form_changes = [
        "Сейчас один флажок объединяет обработку данных, правила парка и программу лояльности. Просим юриста дать отдельные формулировки и указать, какие подтверждения обязательны.",
        "Ссылки сейчас ведут на документы одного парка. В рабочей версии они должны автоматически меняться в зависимости от выбранного парка.",
        "При регистрации через кассира согласия должен давать сам гость. До запуска тестовую SMS-проверку и локальное хранение нужно заменить промышленным решением.",
    ]
    for item in form_changes:
        add_list_item(doc, item, bullet_num)

    add_callout(
        doc,
        "Стоп до согласования.",
        "Не переносить текущий общий флажок и ссылки одного парка во все парки без финальных текстов юриста.",
        fill=PALE_YELLOW,
    )

    doc.add_heading("7. Что юрист должен передать в результате", level=1)
    deliverables = [
        "Короткую таблицу по каждому парку: юрлицо, оператор данных, уведомление регулятора, готовность документов и оставшиеся действия.",
        "Проверенные документы каждого парка и готовые тексты согласий с точными указаниями для формы: флажки, обязательность и ссылки.",
        "Шаблон договора или приложения с платформой и подрядчиками, а также краткий внутренний порядок работы с данными.",
        "Итоговый вывод по каждому парку: можно запускать / можно после выполнения условий / пока нельзя запускать.",
    ]
    _, result_num = add_numbering_definitions(doc)
    for item in deliverables:
        add_list_item(doc, item, result_num)

    doc.add_heading("8. Что нужно предоставить юристу от нас", level=1)
    needed = [
        "Публичную тестовую ссылку на форму регистрации и доступ к RFID-панели.",
        "Список юрлиц всех парков и существующие уведомления Роскомнадзора.",
        "Схему хранения данных: где находится база, кто ее администрирует и какие подрядчики имеют доступ.",
        "Финальный список интеграций: касса, SMS, CRM, СКУД, аналитика, Wallet и другие сервисы.",
    ]
    for item in needed:
        add_list_item(doc, item, bullet_num)

    p = doc.add_paragraph()
    p.add_run("Ссылки: ").bold = True
    p.add_run("пример формы регистрации — ")
    run = p.add_run("[ВСТАВИТЬ ПУБЛИЧНУЮ ТЕСТОВУЮ ССЫЛКУ]")
    set_run_font(run, bold=True, color="9B1C1C")
    p.add_run("; локальный прототип — ")
    run = p.add_run("guest-registration/index.html")
    set_run_font(run, italic=True)
    p.add_run("; документы парков — ")
    add_hyperlink(p, "hello-park.ru/r", "https://hello-park.ru/r")

    add_callout(
        doc,
        "Итог.",
        "После работы юриста у продуктовой команды должен быть простой комплект текстов и правил для интерфейса, а у каждого парка - понятный список обязательных действий до запуска.",
    )

    doc.core_properties.title = "Краткое ТЗ юристу по внедрению RFID-панели Hello Park"
    doc.core_properties.subject = "Регистрация гостей, персональные данные, RFID и программа лояльности"
    doc.core_properties.author = "Hello Park"
    doc.core_properties.comments = "Краткий рабочий документ для юридической подготовки запуска."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
