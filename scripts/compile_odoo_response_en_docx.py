import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    scripts_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(scripts_dir)
    docs_dir = os.path.join(project_dir, 'docs')
    output_path = os.path.join(docs_dir, 'Response_to_Odoo_Team_RFID_Integration.docx')

    doc = docx.Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Base style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(35, 35, 35)

    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(17, 24, 39)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(17, 24, 39)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(31, 41, 55)
        return p

    def add_h4(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(55, 65, 81)
        return p

    def add_callout(text_or_runs, prefix="", bg_color="F9FAFB"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, bg_color)
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if prefix:
            r = p.add_run(prefix)
            r.bold = True
            r.font.color.rgb = RGBColor(17, 24, 39)
        if isinstance(text_or_runs, str):
            p.add_run(text_or_runs)
        elif isinstance(text_or_runs, list):
            for item in text_or_runs:
                if isinstance(item, tuple):
                    r = p.add_run(item[0])
                    if len(item) > 1 and item[1]:
                        r.bold = True
                    if len(item) > 2 and item[2]:
                        r.font.color.rgb = item[2]
                else:
                    p.add_run(str(item))
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, "282C34")
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(235, 235, 235)
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    # --- Title ---
    add_title("Response to Odoo Integration Team: RFID Panel & Hello Park Guest Registration")

    # --- Questions List ---
    add_h2("Questions List")
    p_q_intro = doc.add_paragraph("The integration team requested the following details for the BRD (Business Requirements Document):")
    p_q_intro.paragraph_format.space_after = Pt(4)

    q1 = doc.add_paragraph(style='List Number')
    q1.add_run("RFID Band Binding & Unbinding API Documentation:").bold = True
    q1.paragraph_format.space_after = Pt(1)
    for sub in [
        "API endpoints and authentication details",
        "Request and response specifications",
        "Mandatory parameters and validation rules",
        "Sample request/response payloads",
        "Error codes and handling mechanism",
        "Business rules and dependencies",
        "Sample integration code (if available)"
    ]:
        p_sub = doc.add_paragraph(style='List Bullet 2')
        p_sub.paragraph_format.space_after = Pt(1)
        p_sub.add_run(sub)

    q2 = doc.add_paragraph(style='List Number')
    q2.add_run("Existing Hello Parks Guest Registration Form:").bold = True
    q2.paragraph_format.space_before = Pt(4)
    q2.paragraph_format.space_after = Pt(1)
    for sub in [
        "The guest registration form currently used for the Hello Park setup",
        "Details of all data fields captured during registration",
        "Mandatory/optional fields and validation rules",
        "Consent, terms & conditions, and applicable declarations",
        "Customer identification and mapping logic",
        "UAT credentials for the RFID Cashier Panel and Guest Registration Form"
    ]:
        p_sub = doc.add_paragraph(style='List Bullet 2')
        p_sub.paragraph_format.space_after = Pt(1)
        p_sub.add_run(sub)

    # --- PART 1 ---
    add_h2("PART 1. User Journey & Architecture")

    add_callout(
        [
            ("Interactive Demonstration Environments (Demo mode, no credentials required):\n", True),
            ("• Mobile Guest Registration Form (SPA): ", True),
            ("https://aihrtest-create.github.io/rfid-panel/guest-registration/\n", False),
            ("• RFID Cashier Panel: ", True),
            ("https://aihrtest-create.github.io/rfid-panel/\n\n", False),
            ("(These interfaces are deployed as interactive, live web prototypes. You can test the end-to-end user experience directly in your browser without login credentials).", False)
        ],
        prefix="🔗 ",
        bg_color="FFF7ED"
    )

    add_h3("1.1. High-Level System Architecture (How It Works Currently)")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Guest Registration & Game System: ").bold = True
    p.add_run("Guest registration, child profiles, game progress across interactive attractions, points, and virtual Avatars are hosted and managed on the Hello Park side.")
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Park POS (Cash Desk): ").bold = True
    p.add_run("Ticket sales and receipt printing/fiscalization take place on the park's POS system.")
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Controlled Wristband Issuance: ").bold = True
    p.add_run("A physical RFID wristband is issued to a child only after confirmed ticket payment at the POS. Upon fiscalizing a receipt, the POS sends a payment signal (activation token) to Hello Park, unlocking wristband assignment. Wristbands cannot be assigned without an active paid admission ticket.")
    p.paragraph_format.space_after = Pt(6)

    add_h3("1.2. User Flow Diagram")
    img_flow = os.path.join(docs_dir, 'screenshots', 'user_flow_diagram.png')
    if os.path.exists(img_flow):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        p_img.add_run().add_picture(img_flow, width=Inches(6.6))

    steps = [
        ("Step 1 (Guest): ", "Upon entering the park, the guest scans a QR code using their smartphone and fills out a lightweight mobile web registration form (phone number verified via SMS OTP, parent's full name, children's names and dates of birth)."),
        ("Step 2 (Server): ", "Family profile data is automatically stored in Hello Park's global and local databases and instantly appears in the RFID Cashier Panel web interface on the cashier's computer."),
        ("Step 3 (Park POS): ", "The guest approaches the cash desk and pays for admission tickets. At this point, the cashier cannot bind a wristband yet, as no activation tokens have been granted."),
        ("Step 4 (Integration): ", "Immediately after receipt fiscalization/payment, the park's POS sends a lightweight HTTP POST request (webhook) to the Hello Park backend containing the receipt number and the count of admission tickets."),
        ("Step 5 (Cashier): ", "The RFID Cashier Panel immediately receives the activation credits. The cashier selects the child, taps a clean RFID wristband against the desktop USB reader, and the wristband is bound. The child enters the park to play.")
    ]
    for s_title, s_desc in steps:
        p_step = doc.add_paragraph(style='List Number')
        p_step.paragraph_format.space_after = Pt(3)
        p_step.add_run(s_title).bold = True
        p_step.add_run(s_desc)

    add_h3("1.3. Cashier Workspace (RFID Cashier Panel)")
    p_cash = doc.add_paragraph("A dedicated web application running on the cashier's computer next to the POS terminal. The interface is optimized for rapid cashier operations:")
    p_cash.paragraph_format.space_after = Pt(4)

    img_main = os.path.join(docs_dir, 'screenshots', 'main_screen_en.png')
    if os.path.exists(img_main):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        p_img.add_run().add_picture(img_main, width=Inches(6.6))

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Parent Row: ").bold = True
    p.add_run("Full name, phone number, list of linked children.")
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Child Statuses:").bold = True
    p.paragraph_format.space_after = Pt(1)

    for st_title, st_desc in [
        ("NONE (Grey): ", "Child is registered in the database, but no wristband has been assigned yet."),
        ("BRACELET (Orange): ", "Wristband has been linked at the desk; child is currently active in the park."),
        ("AVATAR (Purple): ", "Child has interacted with game projections in the park. Game progress, level, XP, and prize achievements are persistently saved to this profile.")
    ]:
        p_st = doc.add_paragraph(style='List Bullet 2')
        p_st.paragraph_format.space_after = Pt(1)
        p_st.add_run(st_title).bold = True
        p_st.add_run(st_desc)

    add_h3("1.4. Key Operational Scenarios")

    add_h4("Scenario 1: First-Time Visit (New Family)")
    s1_steps = [
        "Parent fills out the online registration form upon arrival.",
        "Approaches the desk; cashier processes admission tickets on the park's POS.",
        "The POS triggers an payment event → activation credits light up in the Hello Park panel.",
        "Cashier taps a wristband on the desktop reader → the wristband is linked to the child profile → the child enters the attractions."
    ]
    for step in s1_steps:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(step)

    add_h4("Scenario 2: Returning Guest (Progress Continuity)")
    p_desc = doc.add_paragraph("Children frequently return to the park (next week, next month, or months later):")
    p_desc.paragraph_format.space_after = Pt(2)
    s2_steps = [
        "The physical wristband from their previous visit was returned upon exit, but their game progress (Avatar, level, accumulated points, unlocked rewards) remains permanently saved in the Hello Park database.",
        "Upon return, the cashier locates the family in the system (by phone number or name).",
        "The cashier takes a clean physical wristband, clicks \"New bracelet\" next to the child's profile, and taps it on the reader.",
        "The new physical RFID UID is attached to the child's permanent profile (child_id). When the child taps the wristband at any park attraction, their saved Avatar, level, and points are instantly restored."
    ]
    for step in s2_steps:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(step)

    add_h4("Scenario 3: Wristband Return & Same-Day Reuse")
    p_desc = doc.add_paragraph("In high-volume parks, wristbands circulate continuously:")
    p_desc.paragraph_format.space_after = Pt(2)
    s3_steps = [
        "Upon exiting the park, the child drops the wristband at the desk.",
        "The cashier clicks \"Clear bracelet\" in the panel and taps it on the reader — the wristband is unbound in 1 second. The child's game progress remains completely intact.",
        "The cleared wristband is immediately returned to the pool and can be assigned to the next arriving guest.",
        "Automated Nightly Wipe: Every night at 03:00 AM, the Hello Park server automatically resets all active physical RFID wristband associations. In the morning, all wristbands in the park are clean and ready for issuance."
    ]
    for step in s3_steps:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(step)

    # --- PART 2 ---
    add_h2("PART 2. Technical API Specifications (Binding & POS Integration)")

    add_h3("2.1. Architectural Approaches: Who Performs Wristband Binding?")
    doc.add_paragraph("Before finalizing the BRD, we need to align on one of two implementation models:").paragraph_format.space_after = Pt(4)

    add_h4("Option A: Hello Park RFID Panel")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("How it works: ").bold = True
    p.add_run("Cashiers use the standalone Hello Park web panel with a connected desktop USB RFID reader.")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Role of Odoo: ").bold = True
    p.add_run("Odoo does not interact with RFID hardware directly. Odoo only transmits a ticket payment webhook (POST /api/v1/pos/events).")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Result: ").bold = True
    p.add_run("Activation slots are credited to the Hello Park panel, and the cashier assigns wristbands inside the Hello Park UI.")
    p.paragraph_format.space_after = Pt(6)

    add_h4("Option B: Direct Binding from Odoo POS")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("How it works: ").bold = True
    p.add_run("Odoo POS connects to an RFID reader directly within the Odoo POS terminal interface.")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Role of Odoo: ").bold = True
    p.add_run("After scanning the physical wristband, Odoo calls the external Hello Park API to link child_id + rfid_uid.")
    p.paragraph_format.space_after = Pt(6)

    add_h3("2.2. Specification for Option A (Payment Webhook from POS to Hello Park)")
    doc.add_paragraph("The POS sends this notification immediately upon payment/receipt fiscalization whenever the transaction includes at least one admission ticket.").paragraph_format.space_after = Pt(4)

    p_spec = doc.add_paragraph()
    p_spec.add_run("URL: ").bold = True
    p_spec.add_run("POST https://<HP_SERVER_URL>/api/v1/pos/events\n")
    p_spec.add_run("Headers:\n").bold = True
    p_spec.paragraph_format.space_after = Pt(2)

    add_code_block("Content-Type: application/json; charset=utf-8\nAuthorization: Bearer <API_SECRET_KEY>")

    doc.add_paragraph("Request Body (JSON):").bold = True
    add_code_block('{\n  "receipt_number": "REC-2026-0904-001",\n  "count": 2\n}')

    p_params = doc.add_paragraph()
    p_params.add_run("Parameters:\n").bold = True
    p_p1 = doc.add_paragraph(style='List Bullet')
    p_p1.add_run("receipt_number (String, Required): ").bold = True
    p_p1.add_run("Unique receipt/transaction identifier from Odoo. Used for idempotency (prevents duplicate activation credits on network retries).")
    p_p1.paragraph_format.space_after = Pt(1)
    p_p2 = doc.add_paragraph(style='List Bullet')
    p_p2.add_run("count (Integer, Required): ").bold = True
    p_p2.add_run("Total number of paid child admission tickets in the receipt.")
    p_p2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Success Response (200 OK):").bold = True
    add_code_block('{\n  "status": "ok"\n}')

    doc.add_paragraph("Error Codes:").bold = True
    p_e1 = doc.add_paragraph(style='List Bullet')
    p_e1.add_run("401 Unauthorized: ").bold = True
    p_e1.add_run("Invalid or missing secret authorization key (API_SECRET_KEY).")
    p_e1.paragraph_format.space_after = Pt(1)
    p_e2 = doc.add_paragraph(style='List Bullet')
    p_e2.add_run("400 Bad Request: ").bold = True
    p_e2.add_run("Missing mandatory parameters or malformed JSON.")
    p_e2.paragraph_format.space_after = Pt(6)

    add_h3("2.3. Specification for Option B (Direct Bind / Unbind APIs for Odoo)")
    doc.add_paragraph("If Odoo performs the wristband scanning directly on its end:").paragraph_format.space_after = Pt(4)

    add_h4("Method 1: Bind Wristband")
    p = doc.add_paragraph()
    p.add_run("URL: ").bold = True
    p.add_run("POST https://<HP_SERVER_URL>/api/v1/rfid/bind\n")
    p.add_run("Headers: ").bold = True
    p.add_run("Authorization: Bearer <API_SECRET_KEY>, Content-Type: application/json")
    p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("Request Body:").bold = True
    add_code_block('{\n  "child_id": "CH-10023",\n  "rfid_uid": "04A1B2C3D4",\n  "receipt_number": "REC-2026-0904-001"\n}')

    doc.add_paragraph("Success Response (200 OK):").bold = True
    add_code_block('{\n  "status": "success",\n  "message": "Wristband successfully bound",\n  "data": {\n    "child_id": "CH-10023",\n    "rfid_uid": "04A1B2C3D4",\n    "avatar_status": "ready",\n    "bound_at": "2026-09-04T14:30:00Z"\n  }\n}')

    doc.add_paragraph("Error Codes:").bold = True
    for c_err, desc_err in [
        ("400 Bad Request: ", "Missing mandatory parameters."),
        ("404 Not Found: ", "Child with specified child_id not found in registry."),
        ("409 Conflict: ", "This rfid_uid is already assigned to another active child in the park."),
        ("422 Unprocessable Entity: ", "No available ticket activations remain for this receipt number.")
    ]:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(c_err).bold = True
        p.add_run(desc_err)

    add_h4("Method 2: Unbind Wristband (Exit / Same-Day Reuse)")
    p = doc.add_paragraph()
    p.add_run("URL: ").bold = True
    p.add_run("POST https://<HP_SERVER_URL>/api/v1/rfid/unbind\n")
    p.add_run("Headers: ").bold = True
    p.add_run("Authorization: Bearer <API_SECRET_KEY>, Content-Type: application/json")
    p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("Request Body:").bold = True
    add_code_block('{\n  "rfid_uid": "04A1B2C3D4"\n}')

    doc.add_paragraph("Success Response (200 OK):").bold = True
    add_code_block('{\n  "status": "success",\n  "message": "Wristband unbound and released for reuse"\n}')

    add_h3("2.4. Core Business Rules & Dependencies")
    rules = [
        ("Separation of Physical Token and Game Profile: ", "The physical RFID wristband is purely a temporary day token. The child's Avatar, points, levels, and unlocked prizes are permanently tied to child_id in the Hello Park database and are never wiped when a wristband is removed."),
        ("Returning Guests: ", "On subsequent visits, a new physical wristband is issued, but the binding request must pass the same permanent child_id."),
        ("Duplicate Prevention (Conflicts): ", "A single wristband cannot be assigned to two children at the same time. The system returns 409 Conflict if the wristband is already active."),
        ("Automated Daily Reset: ", "Every night, all active physical RFID wristband assignments are cleared, leaving all hardware wristbands ready for reassignment the next morning.")
    ]
    for r_t, r_d in rules:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(r_t).bold = True
        p.add_run(r_d)

    # --- PART 3 ---
    add_h2("PART 3. Guest Registration Form (Hello Park Setup)")

    add_h3("3.1. Registration Form Fields & Structure")
    doc.add_paragraph("The guest registration process is organized into sequential steps optimized for mobile browsers:").paragraph_format.space_after = Pt(3)

    add_h4("Step 1: Contact Information & Legal Consent")
    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run("Phone Number (phone): ").bold = True
    p1.add_run("Mandatory. International phone format (mask and digit count configured according to the park's country). 4-digit SMS OTP verification (simulated in demo mode; accepts any 4 digits except 0000).")
    p1.paragraph_format.space_after = Pt(2)

    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run("Terms & Legal Declarations (Checkbox): ").bold = True
    p2.add_run("Mandatory (SMS sending is disabled until accepted). The guest accepts: 1) Park Visiting Rules; 2) Loyalty Program Terms; 3) Personal Data Processing Policy.")
    p2.paragraph_format.space_after = Pt(4)

    add_h4("Step 2: Family & Children Details")
    f_items = [
        ("Parent Full Name (fio): ", "Mandatory, string (First Name, Last Name, Middle Name)."),
        ("Children List (children): ", "Mandatory, minimum 1 child required. Additional children can be added dynamically."),
        ("• Child First Name (name): ", "Mandatory, text string."),
        ("• Date of Birth (dob): ", "Mandatory, date format (YYYY-MM-DD or localized date picker). Required to calibrate age-appropriate games and quests."),
        ("Optional / Partner-Specific Fields (Configurable per park): ", "\"How did you hear about us?\" (marketing attribution / acquisition channel) and \"City\" (place of residence).")
    ]
    for ft, fd in f_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(ft).bold = True
        p.add_run(fd)

    add_h4("Step 3: Data Persistence & Display (Registration Completion)")
    doc.add_paragraph("Upon form submission:").paragraph_format.space_after = Pt(2)
    for res_item in [
        "Family data is automatically saved to Hello Park's global and local databases.",
        "The parent and child profiles immediately appear in the registered guest queue within the RFID Cashier Panel.",
        "The cashier looks up the family by phone number or name to assign wristbands."
    ]:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(res_item)

    add_callout(
        "For selected parks, an encoded digital loyalty QR card is generated on Step 3. When enabled, the cashier can simply scan this QR code with a 2D desktop scanner to open the family record instantly, avoiding manual phone searches.",
        prefix="Optional Feature (Loyalty QR Code Generation): ",
        bg_color="F3F4F6"
    )

    add_h3("3.2. Customer Identification & Mapping Logic")
    m_items = [
        ("Parent Identifier (Primary Key): ", "Mobile phone number (phone)."),
        ("Child Identifier: ", "Permanent child_id generated upon profile creation."),
        ("Odoo Mapping: ", "The Odoo customer profile should store either the external child_id or the primary family phone number to ensure recurring ticket purchases consistently resolve to the existing child_id.")
    ]
    for mt, md in m_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(mt).bold = True
        p.add_run(md)

    doc.save(output_path)
    print(f"Successfully created: {output_path}")

if __name__ == '__main__':
    create_document()
