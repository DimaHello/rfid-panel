import docx

doc = docx.Document('/Users/dima/Desktop/My project/Rfid-панель/docs/google_doc_downloaded.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.text}")

print("\n--- TABLES ---")
for t_idx, tbl in enumerate(doc.tables):
    print(f"\nTable {t_idx}:")
    for r_idx, row in enumerate(tbl.rows):
        row_text = [cell.text.strip() for cell in row.cells]
        print(f"  Row {r_idx}: {row_text}")
