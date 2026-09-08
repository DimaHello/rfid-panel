import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx():
    scripts_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(scripts_dir)
    docs_dir = os.path.join(project_dir, 'docs')
    output_path = os.path.join(docs_dir, 'ТЗ_Интеграция_Кассы_Активация_Браслетов.docx')

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Интеграция кассового терминала с сервером аватаров HP")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(40, 40, 40)

    def add_section_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_sub_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, "282C34")
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(235, 235, 235)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_callout(text, prefix="Ключевое: ", bg_color="F4F5F9"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, bg_color)
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if prefix:
            r_p = p.add_run(prefix)
            r_p.bold = True
            r_p.font.color.rgb = RGBColor(0, 0, 0)
        p.add_run(text)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Архитектура
    add_section_heading("1. Общая архитектура и пользовательский путь")
    
    add_callout(
        "Функционал регистрации гостей и хранение базы пользователей находится на нашей стороне. "
        "После регистрации система ожидает от кассы сообщение об оплате билета (токен активации), "
        "чтобы разрешить привязку RFID-браслета к ребёнку в нашей игровой системе "
        "(для сохранения игрового прогресса на аттракционах парка).",
        prefix="Ключевое: "
    )

    doc.add_paragraph("Схема пользовательского пути (User Flow):").paragraph_format.space_after = Pt(4)

    img_flow = os.path.join(docs_dir, 'screenshots', 'user_flow_diagram.png')
    if os.path.exists(img_flow):
        p_img_f = doc.add_paragraph()
        p_img_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img_f.paragraph_format.space_before = Pt(4)
        p_img_f.paragraph_format.space_after = Pt(6)
        p_img_f.add_run().add_picture(img_flow, width=Inches(6.6))

    # Шаг 1
    add_sub_heading("Шаг 1. Гость заполняет анкету (на нашей стороне)")
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_after = Pt(3)
    p_s1.add_run("Гость при входе в парк сканирует ")
    p_s1.add_run("QR-код").bold = True
    p_s1.add_run(" с телефона и попадает в мобильную веб-анкету регистрации. В ней он:")
    
    p_s1_1 = doc.add_paragraph(style='List Bullet')
    p_s1_1.paragraph_format.space_after = Pt(1)
    p_s1_1.add_run("Вводит номер телефона и подтверждает его по СМС.")
    
    p_s1_2 = doc.add_paragraph(style='List Bullet')
    p_s1_2.paragraph_format.space_after = Pt(1)
    p_s1_2.add_run("Соглашается с правилами парка и программой лояльности.")
    
    p_s1_3 = doc.add_paragraph(style='List Bullet')
    p_s1_3.paragraph_format.space_after = Pt(6)
    p_s1_3.add_run("Указывает ФИО родителя и добавляет данные детей (имя, дата рождения).")

    # Шаг 2
    add_sub_heading("Шаг 2. Данные попадают в нашу базу и RFID-панель")
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_after = Pt(3)
    p_s2.add_run("После заполнения анкеты данные семьи автоматически сохраняются в ")
    p_s2.add_run("нашей базе данных").bold = True
    p_s2.add_run(" и мгновенно появляются в веб-интерфейсе ")
    p_s2.add_run("RFID-панели кассира").bold = True
    p_s2.add_run(".")

    p_s2_desc = doc.add_paragraph()
    p_s2_desc.paragraph_format.space_after = Pt(3)
    p_s2_desc.add_run("RFID-панель кассира — это веб-приложение, открытое на компьютере кассира в парке. В ней отображается таблица зарегистрированных аккаунтов:")
    
    p_s2_l1 = doc.add_paragraph(style='List Bullet')
    p_s2_l1.paragraph_format.space_after = Pt(1)
    p_s2_l1.add_run("Аккаунт родителя: ").bold = True
    p_s2_l1.add_run("ФИО, номер телефона.")

    p_s2_l2 = doc.add_paragraph(style='List Bullet')
    p_s2_l2.paragraph_format.space_after = Pt(4)
    p_s2_l2.add_run("Привязанные дети: ").bold = True
    p_s2_l2.add_run("имя, дата рождения, статус браслета (есть / нет).")

    img_panel_en = os.path.join(docs_dir, 'screenshots', 'main_screen_en.png')
    if os.path.exists(img_panel_en):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_panel_en, width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run("Рис. 1 — RFID Cashier Panel (таблица зарегистрированных аккаунтов)")
        r_cap.italic = True
        r_cap.font.size = Pt(9)
        r_cap.font.color.rgb = RGBColor(120, 120, 120)

    # Шаг 3
    add_sub_heading("Шаг 3. Гость оплачивает билет на кассе")
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_after = Pt(3)
    p_s3.add_run("Гость подходит к кассе и оплачивает входные билеты. На этом этапе кассир ")
    p_s3.add_run("не может").bold = True
    p_s3.add_run(" привязать браслет ребёнку — у него нет токена активации.")

    p_s3_2 = doc.add_paragraph()
    p_s3_2.paragraph_format.space_after = Pt(4)
    p_s3_2.add_run("Токен активации — это разрешение на привязку браслета. Он появляется в RFID-панели ")
    p_s3_2.add_run("только после того, как касса подтвердит оплату").bold = True
    p_s3_2.add_run(", отправив запрос на наш бэкенд. Без токена активации кнопка привязки браслета в панели заблокирована. Это гарантирует, что браслеты выдаются только оплатившим гостям.")

    # Шаг 4
    add_sub_heading("Шаг 4. Касса отправляет запрос на наш бэкенд (Интеграция)")
    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_after = Pt(4)
    p_s4.add_run("Это единственный шаг, который требует доработки на стороне кассового ПО. После фискализации чека касса отправляет ")
    p_s4.add_run("один простой HTTP POST запрос").bold = True
    p_s4.add_run(" на наш сервер (спецификация приведена в разделе 2.1 ниже).")

    # Шаг 5
    add_sub_heading("Шаг 5. Кассир привязывает браслет к ребёнку")
    p_s5 = doc.add_paragraph()
    p_s5.paragraph_format.space_after = Pt(3)
    p_s5.add_run("После получения запроса от кассы в RFID-панели ")
    p_s5.add_run("моментально появляется токен активации").bold = True
    p_s5.add_run(" (счетчик доступных браслетов). Кассир:")
    
    p_s5_1 = doc.add_paragraph(style='List Bullet')
    p_s5_1.paragraph_format.space_after = Pt(1)
    p_s5_1.add_run("Выбирает ребёнка из списка.")

    p_s5_2 = doc.add_paragraph(style='List Bullet')
    p_s5_2.paragraph_format.space_after = Pt(1)
    p_s5_2.add_run("Прикладывает чистый RFID-браслет к настольному считывателю.")

    p_s5_3 = doc.add_paragraph(style='List Bullet')
    p_s5_3.paragraph_format.space_after = Pt(6)
    p_s5_3.add_run("Браслет привязывается к профилю ребёнка — ребёнок может идти играть.")

    # 2. Спецификации запросов
    add_section_heading("2. Спецификации запросов")

    # 2.1 Получение данных от кассы об оплате
    add_sub_heading("2.1 Получение данных от кассы об оплате")
    p_pos_intro = doc.add_paragraph()
    p_pos_intro.paragraph_format.space_after = Pt(3)
    p_pos_intro.add_run("Поставщик кассы должен реализовать отправку сигнала об оплате билета на наш локальный сервер.")

    p_trig1 = doc.add_paragraph(style='List Bullet')
    p_trig1.paragraph_format.space_after = Pt(2)
    p_trig1.add_run("Момент отправки: ").bold = True
    p_trig1.add_run("Сразу после успешной оплаты и фискализации/печати чека на кассе.")

    p_trig2 = doc.add_paragraph(style='List Bullet')
    p_trig2.paragraph_format.space_after = Pt(4)
    p_trig2.add_run("Условие отправки: ").bold = True
    p_trig2.add_run("Если в чеке присутствует хотя бы 1 входной билет (товар из категории «Билеты»). Если в чеке только бар, кафе или сувениры — запрос не отправляется.")

    p_req_t = doc.add_paragraph()
    p_req_t.paragraph_format.space_after = Pt(2)
    p_req_t.add_run("Запрос:").bold = True

    p_req1 = doc.add_paragraph(style='List Bullet')
    p_req1.paragraph_format.space_after = Pt(2)
    p_req1.add_run("Метод: ").bold = True
    p_req1.add_run("POST")

    p_req2 = doc.add_paragraph(style='List Bullet')
    p_req2.paragraph_format.space_after = Pt(2)
    p_req2.add_run("Адрес (URL): ").bold = True
    p_req2.add_run("https://<HP_LOCAL_SERVER_URL>/api/v1/pos/events")

    p_req3 = doc.add_paragraph(style='List Bullet')
    p_req3.paragraph_format.space_after = Pt(2)
    p_req3.add_run("Заголовки (Headers): ").bold = True

    p_h1 = doc.add_paragraph(style='List Bullet 2')
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.add_run("Content-Type: application/json; charset=utf-8")

    p_h2 = doc.add_paragraph(style='List Bullet 2')
    p_h2.paragraph_format.space_after = Pt(4)
    r_auth = p_h2.add_run("Authorization: Bearer <API_SECRET_KEY>")
    r_auth.bold = True
    p_h2.add_run(" — секретный ключ авторизации кассы (выдаётся нами при подключении). Без валидного ключа сервер отклоняет запрос с ошибкой 401 Unauthorized.")

    doc.add_paragraph("Тело запроса (JSON Payload):").paragraph_format.space_after = Pt(3)
    json_example = """{
  "receipt_number": "12345",
  "count": 2
}"""
    add_code_block(json_example)

    doc.add_paragraph("Описание параметров JSON:").paragraph_format.space_after = Pt(4)

    # Table of fields
    table_fields = doc.add_table(rows=3, cols=4)
    table_fields.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Параметр", "Тип", "Обязателен", "Описание и пример"]
    col_widths = [Inches(1.6), Inches(0.9), Inches(1.1), Inches(3.2)]

    for i, h in enumerate(headers):
        cell = table_fields.rows[0].cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "282C34")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    rows_data = [
        ("receipt_number", "String", "Да", "Номер пробитого чека (например: \"12345\" или \"00123\")"),
        ("count", "Integer", "Да", "Количество оплаченных билетов в чеке (например: 1, 2, 3)")
    ]

    for row_idx, data in enumerate(rows_data):
        row = table_fields.rows[row_idx + 1]
        bg = "FFFFFF" if row_idx % 2 == 0 else "F9FAFB"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if col_idx == 0:
                run.font.name = 'Consolas'
                run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Server Response
    p_resp_t = doc.add_paragraph()
    p_resp_t.paragraph_format.space_after = Pt(2)
    p_resp_t.add_run("Формат ответа сервера:").bold = True
    
    doc.add_paragraph("При успешном приеме события бэкенд возвращает статус 200 OK:").paragraph_format.space_after = Pt(3)
    add_code_block('{\n  "status": "ok"\n}')

    p_err = doc.add_paragraph()
    p_err.paragraph_format.space_after = Pt(4)
    p_err.add_run("Если токен авторизации отсутствует или указан неверно, сервер вернет ")
    p_err.add_run("401 Unauthorized").bold = True
    p_err.add_run(" и отклонит операцию.")

    # Fault tolerance
    p_ft_t = doc.add_paragraph()
    p_ft_t.paragraph_format.space_after = Pt(2)
    p_ft_t.add_run("Отказоустойчивость и повторные отправки:").bold = True

    p_idemp = doc.add_paragraph(style='List Bullet')
    p_idemp.paragraph_format.space_after = Pt(2)
    p_idemp.add_run("Защита от дублей (идемпотентность): ").bold = True
    p_idemp.add_run("Идентификация операции происходит по полю receipt_number. Если из-за сбоя сети касса отправит повторный запрос по тому же номеру чека, система не начислит дублирующие браслеты, а повторно вернет 200 OK.")

    p_off = doc.add_paragraph(style='List Bullet')
    p_off.paragraph_format.space_after = Pt(6)
    p_off.add_run("Офлайн-очередь на кассе: ").bold = True
    p_off.add_run("Если в момент закрытия чека сеть временно недоступна, кассовый модуль должен сохранить запрос в локальную очередь и автоматически отправить его при появлении связи.")

    # 2.2 Авторизация на кассе
    add_sub_heading("2.2 Авторизация на кассе")
    p_auth_desc = doc.add_paragraph()
    p_auth_desc.paragraph_format.space_after = Pt(3)
    p_auth_desc.add_run("Если кассе нужно получить данные о пользователе, она отправляет запрос в облачный сервер:")

    add_code_block("GET https://<HP_SERVER>/api/user")

    p_param_desc = doc.add_paragraph()
    p_param_desc.paragraph_format.space_after = Pt(2)
    p_param_desc.add_run("В запросе передается один из параметров:")

    p_p1 = doc.add_paragraph(style='List Bullet')
    p_p1.paragraph_format.space_after = Pt(1)
    p_p1.add_run("Id пользователя (считывается по QR-коду гостя в конце регистрации)")

    p_p2 = doc.add_paragraph(style='List Bullet')
    p_p2.paragraph_format.space_after = Pt(3)
    p_p2.add_run("Номер телефона")

    p_resp_desc = doc.add_paragraph()
    p_resp_desc.paragraph_format.space_after = Pt(6)
    p_resp_desc.add_run("В ответ наше облако отдает профиль пользователя и привязанных детей.")

    # 2.3 Калитки и СКУД
    add_sub_heading("2.3 Калитки и система учета времени (СКУД)")
    p_turn_desc = doc.add_paragraph()
    p_turn_desc.paragraph_format.space_after = Pt(3)
    p_turn_desc.add_run("После присваивания браслета наша система может отправить запрос в систему СКУД (турникеты):")

    p_post_params = doc.add_paragraph()
    p_post_params.paragraph_format.space_after = Pt(2)
    p_post_params.add_run("POST запрос с параметрами:").bold = True

    p_t1 = doc.add_paragraph(style='List Bullet')
    p_t1.paragraph_format.space_after = Pt(1)
    p_t1.add_run("user_id — Id пользователя / ребёнка")

    p_t2 = doc.add_paragraph(style='List Bullet')
    p_t2.paragraph_format.space_after = Pt(1)
    p_t2.add_run("rfid — Номер RFID-браслета")

    p_t3 = doc.add_paragraph(style='List Bullet')
    p_t3.paragraph_format.space_after = Pt(6)
    p_t3.add_run("date — Дата и время регистрации / привязки")

    # 2.4 Вебхук регистрации
    add_sub_heading("2.4 Вебхук регистрации нового пользователя (CRM)")
    p_crm_desc = doc.add_paragraph()
    p_crm_desc.paragraph_format.space_after = Pt(3)
    p_crm_desc.add_run("При регистрации нового пользователя наш сервер может вызывать вебхук сторонней CRM системы и передать ему данные пользователя:")

    p_c1 = doc.add_paragraph(style='List Bullet')
    p_c1.paragraph_format.space_after = Pt(1)
    p_c1.add_run("Id в нашей системе (user_id)")

    p_c2 = doc.add_paragraph(style='List Bullet')
    p_c2.paragraph_format.space_after = Pt(1)
    p_c2.add_run("Контакты (номер телефона)")

    p_c3 = doc.add_paragraph(style='List Bullet')
    p_c3.paragraph_format.space_after = Pt(6)
    p_c3.add_run("Данные профиля (ФИО, пол, возраст, дата рождения)")

    doc.save(output_path)
    print(f"Russian DOCX successfully compiled at: {output_path}")

if __name__ == '__main__':
    create_docx()
