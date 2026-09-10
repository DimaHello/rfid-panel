import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def main():
    scripts_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(scripts_dir)
    docs_dir = os.path.join(project_dir, 'docs')
    screenshots_dir = os.path.join(docs_dir, 'screenshots', 'guest_registration')

    # 1. Full Markdown content
    md_content = """# Техническое задание: Шаг опроса в мобильной анкете гостя (Источник и Город)

Техническое задание на доработку и внедрение обязательного экрана сбора маркетинговых данных и географии гостей в мобильную анкету саморегистрации сети парков **Hello Park**.

---

## 1. Общие сведения и бизнес-цель

* **Модуль:** Мобильная веб-анкета гостя (`/guest-registration`).
* **Интерактивный веб-прототип:**
  * Локальный запуск: [http://127.0.0.1:3333/guest-registration/index.html](http://127.0.0.1:3333/guest-registration/index.html)
  * Репозиторий GitHub: [https://github.com/aihrtest-create/rfid-panel/tree/main/guest-registration](https://github.com/aihrtest-create/rfid-panel/tree/main/guest-registration)
* **Бизнес-цели:**
  1. **Сбор источников привлечения:** определение эффективности рекламных каналов (баннеры, ТЦ, соцсети, блогеры, сарафанное радио).
  2. **География посетителей:** фиксация города проживания для оценки доли местных жителей района, гостей из других округов и туристов.
  3. **Универсальность:** исключение ручной настройки под каждый парк сети (отказ от фиксированных плашек городов парка в пользу умного поиска по общей базе).
  4. **Высокая скорость прохождения:** заполнение занимает до 10–15 секунд благодаря быстрым чипсам и моментальному автокомплиту без зависимости от внешних API.

---

## 2. Место в пользовательском пути (CJM)

Анкета гостя расширена с 3 до **4 последовательных шагов**:

1. **Шаг 1 (Согласие):** Ввод номера телефона, подтверждение юридических согласий парка, верификация через 4-значный СМС-код.
2. **Шаг 2 (Данные семьи):** Ввод ФИО родителя и добавление детей (имя и дата рождения). Кнопка «Далее».
3. **Шаг 3 (Опрос — Новый обязательный шаг):** «Ещё пару вопросов — и в парк! 🎉». Ответ на 2 обязательных вопроса: источник и город. Кнопка «Зарегистрироваться».
4. **Шаг 4 (QR-код):** Экран успешной регистрации с динамическим QR-кодом для считывания на кассе и кнопками Apple / Google Wallet.

---

## 3. Визуальный интерфейс экрана опроса

### Общий вид экрана (Шаг 3)
![Экран опроса в анкете гостя](screenshots/guest_registration/step3_survey_clean.png)

### Работа умного автокомплита города
![Умный автокомплит города с подсказками](screenshots/guest_registration/step3_autocomplete.png)

---

## 4. Спецификация экранных элементов и логика

### 4.1. Шапка экрана и персонаж
* **Прогресс-бар:** Плашка «Шаг 3 из 4: Опрос» с 4 сегментами (первые три активны, цвет `#FF6022`).
* **Заголовок:** Крупный акцентный заголовок: *«Ещё пару вопросов — и в парк!»*.
* **Персонаж:** 3D-маскот Лис в худи и защитных очках, создающий дружелюбный игровой тон.
* **Пояснение:** *«Ответьте на 2 вопроса, чтобы помочь нам сделать Hello Park ещё лучше для вас и ваших детей.»*.

---

### 4.2. Вопрос 1: «Откуда вы узнали о нас?»

* **Заголовок:** `Откуда вы узнали о нас? *` (стиль `.survey-question-title`, 1.05rem / 17px, начертание 800).
* **Быстрые чипсы (топ-4 ответа в 1 тап):**
  1. 🏢 **Увидели в ТЦ**
  2. 📱 **Инстаграм**
  3. 👥 **От знакомых**
  4. 🎈 **Были в парке**
* **Выпадающий список (все 12 вариантов):**
  * Стилизованный селектор с чистой векторной стрелочкой справа:
    * *от знакомых, сайт, инстаграм, интернет, увидели в ТЦ, экраны/баннеры в городе, реклама на/в транспорте, блогер, были в парке, постоянный клиент, не помню, другое.*
* **Синхронизация:** Выбор чипса переключает выпадающий список; выбор из списка подсвечивает соответствующий чипс.
* **Вариант «Другое»:** Плавно раскрывает текстовое поле `[ Уточните, пожалуйста, откуда... ]`.
* **Валидация:** Поле обязательное. При попытке перехода без выбора селектор подсвечивается ошибкой.

---

### 4.3. Вопрос 2: «Из какого вы города?»

* **Заголовок:** `Из какого вы города? *` (стиль `.survey-question-title`, 1.05rem / 17px, начертание 800).
* **Универсальность для всех парков:**
  * **Отсутствуют фиксированные плашки городов парка.** Это исключает необходимость ручной настройки под каждый парк сети.
* **Поле ввода:**
  * Иконка геолокации (пин) слева, плейсхолдер *«Начните вводить город...»*, кнопка очистки (✕) справа.
* **Автономная база (`cities.js`):**
  * Встроен локальный датасет на **1134 города РФ**, отсортированных по численности населения.
  * Работает на клиенте со скоростью **0 мс задержки**, без платных внешних API и без сбоев при слабом интернете в ТЦ.
* **Умный автокомплит (Live Search):**
  * Мгновенный поиск при вводе от 1 символа (топ-10 совпадений).
  * Подсветка совпавших букв оранжевым цветом.
  * Отображение субъекта/области РФ (например: **Московский** *(Москва)*, **Новомосковск** *(Тульская область)*).
  * Свободный ввод: не блокирует отправку, если гость ввёл редкий посёлок или иностранный город.
* **Валидация:** Обязательное поле (минимум 2 символа).

---

### 4.4. Навигация
* **Кнопка «Назад»:** возврат на Шаг 2 (Данные семьи) без потери данных.
* **Кнопка «Зарегистрироваться»:** проверяет поля опроса, генерирует QR-код лояльности и переводит на Шаг 4.

---

## 5. Формат данных для кассовой системы (Payload QR-кода)

```json
{
  "v": 1,
  "fio": "Иванова Мария Сергеевна",
  "phone": "+79991234567",
  "children": [
    { "name": "Александр", "dob": "2018-05-12" }
  ],
  "source": "увидели в ТЦ",
  "sourceOther": "",
  "city": "Самара",
  "ts": 1725965400000
}
```

Строка кодируется в Base64 с префиксом `HPARK:` и сохраняется в QR-коде.
"""
    with open(os.path.join(docs_dir, 'ТЗ_Опрос_в_анкете_регистрации.md'), 'w', encoding='utf-8') as f:
        f.write(md_content.strip() + '\n')
    print('Updated docs/ТЗ_Опрос_в_анкете_регистрации.md')

    # 2. Build Word document (.docx)
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run('ТЕХНИЧЕСКОЕ ЗАДАНИЕ')
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 96, 34)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run('Шаг опроса в мобильной анкете гостя: источник привлечения и город')
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(50, 50, 50)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 96, 34)

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(40, 40, 40)

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(bold_prefix)
        r1.font.bold = True
        p.add_run(text)

    add_heading_1('1. Общие сведения и бизнес-цель')
    add_bullet('Модуль: ', 'Мобильная веб-анкета саморегистрации гостя (/guest-registration).')
    add_bullet('Интерактивный прототип: ', 'http://127.0.0.1:3333/guest-registration/index.html')
    add_bullet('Репозиторий GitHub: ', 'https://github.com/aihrtest-create/rfid-panel/tree/main/guest-registration')
    add_bullet('Бизнес-цели: ', 'Сбор источников привлечения гостей (эффективность рекламы, ТЦ, соцсетей, блогеров) и географии посетителей для сквозной маркетинговой аналитики.')
    add_bullet('Универсальность: ', 'Анкета едина для всей сети парков Hello Park и не требует ручной настройки городов под конкретные филиалы.')

    add_heading_1('2. Место в пользовательском пути (CJM)')
    p_cjm = doc.add_paragraph('Анкета гостя расширена с 3 до 4 шагов:')
    p_cjm.paragraph_format.space_after = Pt(4)
    add_bullet('Шаг 1 (Согласие): ', 'Номер телефона, юридические согласия и СМС-верификация.')
    add_bullet('Шаг 2 (Данные семьи): ', 'ФИО родителя и добавление детей (имя и дата рождения).')
    add_bullet('Шаг 3 (Опрос — Новый обязательный шаг): ', '«Ещё пару вопросов — и в парк! 🎉». Выбор источника и ввод города.')
    add_bullet('Шаг 4 (QR-код): ', 'Готовая карта лояльности для сканирования на кассе парка.')

    add_heading_1('3. Визуальный интерфейс экрана опроса')
    p_img1 = doc.add_paragraph('Рис. 1: Экран опроса (Шаг 3) без фиксированных плашек городов:')
    p_img1.paragraph_format.space_after = Pt(4)
    img1_path = os.path.join(screenshots_dir, 'step3_survey_clean.png')
    if os.path.exists(img1_path):
        doc.add_picture(img1_path, width=Inches(3.2))

    p_img2 = doc.add_paragraph('Рис. 2: Работа умного автокомплита города с подсказками из базы:')
    p_img2.paragraph_format.space_before = Pt(8)
    p_img2.paragraph_format.space_after = Pt(4)
    img2_path = os.path.join(screenshots_dir, 'step3_autocomplete.png')
    if os.path.exists(img2_path):
        doc.add_picture(img2_path, width=Inches(3.2))

    add_heading_1('4. Спецификация полей и логика работы')

    add_heading_2('4.1. Вопрос «Откуда вы узнали о нас?»')
    add_bullet('Быстрые чипсы (топ-4 в 1 тап): ', '«Увидели в ТЦ», «Инстаграм», «От знакомых», «Были в парке».')
    add_bullet('Полный список (12 вариантов): ', 'от знакомых, сайт, инстаграм, интернет, увидели в ТЦ, экраны/баннеры в городе, реклама на/в транспорте, блогер, были в парке, постоянный клиент, не помню, другое.')
    add_bullet('Вариант «Другое»: ', 'Раскрывает текстовую строку для свободного уточнения.')
    add_bullet('Валидация: ', 'Обязательное поле. При отсутствии выбора подсвечивается ошибкой.')

    add_heading_2('4.2. Вопрос «Из какого вы города?»')
    add_bullet('Отказ от локальных плашек: ', 'Плашки конкретных городов удалены, чтобы исключить необходимость настраивать анкету под каждый филиал сети.')
    add_bullet('Автономная база данных: ', 'Встроен датасет cities.js на 1134 города РФ (0 мс задержки, offline-first, без платных внешних API).')
    add_bullet('Умный автокомплит: ', 'Поиск от 1 символа, подсветка совпадений оранжевым цветом, вывод субъекта РФ (области).')
    add_bullet('Свободный ввод: ', 'Позволяет сохранить посёлок, деревню или зарубежный город вне справочника.')

    add_heading_1('5. Формат данных для кассы (Payload QR-кода)')
    p_json = doc.add_paragraph('Данные опроса упаковываются в JSON и кодируются в Base64 с префиксом HPARK: для кассы:')
    p_json.paragraph_format.space_after = Pt(4)

    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    hdr_cells[0].text = 'Поле'
    hdr_cells[1].text = 'Описание / Пример'
    set_cell_background(hdr_cells[0], 'F2F2F2')
    set_cell_background(hdr_cells[1], 'F2F2F2')
    hdr_cells[0].paragraphs[0].runs[0].font.bold = True
    hdr_cells[1].paragraphs[0].runs[0].font.bold = True

    data_rows = [
        ('source', 'Выбранный источник (например, "увидели в ТЦ")'),
        ('sourceOther', 'Уточнение, если выбрано "другое"'),
        ('city', 'Город проживания гостя (например, "Самара")'),
        ('fio', 'ФИО родителя'),
        ('phone', 'Верифицированный номер телефона (+7...)'),
        ('children', 'Массив объектов детей [ { name, dob } ]')
    ]

    for field, desc in data_rows:
        row = tbl.add_row()
        c1, c2 = row.cells
        c1.text = field
        c2.text = desc
        c1.paragraphs[0].runs[0].font.name = 'Consolas'
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_margins(c1, 80, 80, 100, 100)
        set_cell_margins(c2, 80, 80, 100, 100)

    docx_path = os.path.join(docs_dir, 'ТЗ_Опрос_в_анкете_регистрации.docx')
    doc.save(docx_path)
    print('Updated docs/ТЗ_Опрос_в_анкете_регистрации.docx')

if __name__ == '__main__':
    main()
