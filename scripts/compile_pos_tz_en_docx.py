import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx():
    scripts_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(scripts_dir)
    docs_dir = os.path.join(project_dir, 'docs')
    output_path = os.path.join(docs_dir, 'Technical_Specification_POS_Integration_RFID_Wristband_Activation.docx')

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("TECHNICAL SPECIFICATION")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("POS Terminal Integration with HP Avatar Server")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(40, 40, 40)

    def add_section_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_sub_heading(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, "282C34")
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(235, 235, 235)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_callout(text, prefix="Key Concept: ", bg_color="F4F5F9"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, bg_color)
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if prefix:
            r_p = p.add_run(prefix)
            r_p.bold = True
            r_p.font.color.rgb = RGBColor(0, 0, 0)
        p.add_run(text)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. System Architecture & User Journey
    add_section_heading("1. System Architecture and User Journey")
    
    add_callout(
        "Guest registration and user database management are hosted on our side. "
        "After registration, the system waits for a ticket sale event (activation token) from the POS terminal "
        "to enable RFID wristband binding to a child profile in our game system "
        "(to preserve game progress across interactive park attractions).",
        prefix="Key Concept: "
    )

    doc.add_paragraph("User Flow Diagram:").paragraph_format.space_after = Pt(4)

    img_flow = os.path.join(docs_dir, 'screenshots', 'user_flow_diagram.png')
    if os.path.exists(img_flow):
        p_img_f = doc.add_paragraph()
        p_img_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img_f.paragraph_format.space_before = Pt(4)
        p_img_f.paragraph_format.space_after = Pt(6)
        p_img_f.add_run().add_picture(img_flow, width=Inches(6.6))

    # Step 1
    add_sub_heading("Step 1. Guest fills out registration form (on our side)")
    p_s1 = doc.add_paragraph()
    p_s1.paragraph_format.space_after = Pt(3)
    p_s1.add_run("Upon entering the park, the guest scans a ")
    p_s1.add_run("QR code").bold = True
    p_s1.add_run(" with their smartphone and accesses a mobile web registration form. In the form, the guest:")
    
    p_s1_1 = doc.add_paragraph(style='List Bullet')
    p_s1_1.paragraph_format.space_after = Pt(1)
    p_s1_1.add_run("Enters phone number and verifies it via SMS.")
    
    p_s1_2 = doc.add_paragraph(style='List Bullet')
    p_s1_2.paragraph_format.space_after = Pt(1)
    p_s1_2.add_run("Agrees to park rules and loyalty program terms.")
    
    p_s1_3 = doc.add_paragraph(style='List Bullet')
    p_s1_3.paragraph_format.space_after = Pt(6)
    p_s1_3.add_run("Enters parent's full name and adds children details (name, date of birth).")

    # Step 2
    add_sub_heading("Step 2. Data is stored in our database and appears in RFID Panel")
    p_s2 = doc.add_paragraph()
    p_s2.paragraph_format.space_after = Pt(3)
    p_s2.add_run("Upon form submission, family data is automatically saved to ")
    p_s2.add_run("our database").bold = True
    p_s2.add_run(" and instantly appears in the web interface of the ")
    p_s2.add_run("RFID Cashier Panel").bold = True
    p_s2.add_run(".")

    p_s2_desc = doc.add_paragraph()
    p_s2_desc.paragraph_format.space_after = Pt(3)
    p_s2_desc.add_run("The RFID Cashier Panel is a web application running on the cashier's computer at the park. It displays a list of registered guest accounts:")
    
    p_s2_l1 = doc.add_paragraph(style='List Bullet')
    p_s2_l1.paragraph_format.space_after = Pt(1)
    p_s2_l1.add_run("Parent Account: ").bold = True
    p_s2_l1.add_run("Full name, phone number.")

    p_s2_l2 = doc.add_paragraph(style='List Bullet')
    p_s2_l2.paragraph_format.space_after = Pt(4)
    p_s2_l2.add_run("Linked Children: ").bold = True
    p_s2_l2.add_run("Name, date of birth, wristband status (Avatar / Bracelet / None).")

    img_panel_en = os.path.join(docs_dir, 'screenshots', 'main_screen_en.png')
    if os.path.exists(img_panel_en):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.add_run().add_picture(img_panel_en, width=Inches(6.0))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run("Fig. 1 — RFID Cashier Panel (registered accounts table)")
        r_cap.italic = True
        r_cap.font.size = Pt(9)
        r_cap.font.color.rgb = RGBColor(120, 120, 120)

    # Step 3
    add_sub_heading("Step 3. Guest pays for ticket at POS")
    p_s3 = doc.add_paragraph()
    p_s3.paragraph_format.space_after = Pt(3)
    p_s3.add_run("The guest approaches the ticket desk and pays for admission tickets. At this stage, the cashier ")
    p_s3.add_run("cannot").bold = True
    p_s3.add_run(" bind a wristband to a child because no activation token is available yet.")

    p_s3_2 = doc.add_paragraph()
    p_s3_2.paragraph_format.space_after = Pt(4)
    p_s3_2.add_run("The activation token is an authorization grant for wristband binding. It appears in the RFID panel ")
    p_s3_2.add_run("only after the POS terminal confirms the payment").bold = True
    p_s3_2.add_run(" by sending a request to our backend. Without an activation token, the wristband binding action is locked. This ensures wristbands are only issued to paying guests.")

    # Step 4
    add_sub_heading("Step 4. POS terminal sends request to our backend (Integration)")
    p_s4 = doc.add_paragraph()
    p_s4.paragraph_format.space_after = Pt(4)
    p_s4.add_run("This is the only step requiring implementation on the POS software side. Immediately after receipt fiscalization/printing, the POS terminal sends a ")
    p_s4.add_run("single simple HTTP POST request").bold = True
    p_s4.add_run(" to our server (see specification in Section 2.1 below).")

    # Step 5
    add_sub_heading("Step 5. Cashier binds wristband to child")
    p_s5 = doc.add_paragraph()
    p_s5.paragraph_format.space_after = Pt(3)
    p_s5.add_run("Upon receiving the event from the POS terminal, the ")
    p_s5.add_run("activation token appears instantly").bold = True
    p_s5.add_run(" in the RFID panel (available wristband counter increments). The cashier:")
    
    p_s5_1 = doc.add_paragraph(style='List Bullet')
    p_s5_1.paragraph_format.space_after = Pt(1)
    p_s5_1.add_run("Selects the child from the list.")

    p_s5_2 = doc.add_paragraph(style='List Bullet')
    p_s5_2.paragraph_format.space_after = Pt(1)
    p_s5_2.add_run("Taps a clean RFID wristband against the desktop USB/RFID reader.")

    p_s5_3 = doc.add_paragraph(style='List Bullet')
    p_s5_3.paragraph_format.space_after = Pt(6)
    p_s5_3.add_run("The wristband is linked to the child's profile — the child can now enter and play.")

    # 2. Request Specifications
    add_section_heading("2. API Request Specifications")

    # 2.1 Receiving Payment Data from POS Terminal
    add_sub_heading("2.1 Receiving Payment Data from POS Terminal")
    p_pos_intro = doc.add_paragraph()
    p_pos_intro.paragraph_format.space_after = Pt(3)
    p_pos_intro.add_run("The POS vendor must implement sending a ticket payment signal to our local server.")

    p_trig1 = doc.add_paragraph(style='List Bullet')
    p_trig1.paragraph_format.space_after = Pt(2)
    p_trig1.add_run("When to send: ").bold = True
    p_trig1.add_run("Immediately upon successful payment and receipt printing/fiscalization on the POS terminal.")

    p_trig2 = doc.add_paragraph(style='List Bullet')
    p_trig2.paragraph_format.space_after = Pt(4)
    p_trig2.add_run("Trigger condition: ").bold = True
    p_trig2.add_run("When the receipt contains at least 1 admission ticket (item from the 'Tickets' category). If the receipt contains only bar/café items or merchandise, no request is sent.")

    p_req_t = doc.add_paragraph()
    p_req_t.paragraph_format.space_after = Pt(2)
    p_req_t.add_run("Request:").bold = True

    p_req1 = doc.add_paragraph(style='List Bullet')
    p_req1.paragraph_format.space_after = Pt(2)
    p_req1.add_run("Method: ").bold = True
    p_req1.add_run("POST")

    p_req2 = doc.add_paragraph(style='List Bullet')
    p_req2.paragraph_format.space_after = Pt(2)
    p_req2.add_run("URL: ").bold = True
    p_req2.add_run("https://<HP_LOCAL_SERVER_URL>/api/v1/pos/events")

    p_req3 = doc.add_paragraph(style='List Bullet')
    p_req3.paragraph_format.space_after = Pt(2)
    p_req3.add_run("Headers: ").bold = True

    p_h1 = doc.add_paragraph(style='List Bullet 2')
    p_h1.paragraph_format.space_after = Pt(2)
    p_h1.add_run("Content-Type: application/json; charset=utf-8")

    p_h2 = doc.add_paragraph(style='List Bullet 2')
    p_h2.paragraph_format.space_after = Pt(4)
    r_auth = p_h2.add_run("Authorization: Bearer <API_SECRET_KEY>")
    r_auth.bold = True
    p_h2.add_run(" — secret authorization token for the POS (provided by us upon connection). Without a valid key, the server rejects the request with HTTP 401 Unauthorized.")

    doc.add_paragraph("JSON Payload:").paragraph_format.space_after = Pt(3)
    json_example = """{
  "receipt_number": "12345",
  "count": 2
}"""
    add_code_block(json_example)

    doc.add_paragraph("Parameters Description:").paragraph_format.space_after = Pt(4)

    # Table of fields
    table_fields = doc.add_table(rows=3, cols=4)
    table_fields.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Parameter", "Type", "Required", "Description & Example"]
    col_widths = [Inches(1.6), Inches(0.9), Inches(1.1), Inches(3.2)]

    for i, h in enumerate(headers):
        cell = table_fields.rows[0].cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "282C34")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    rows_data = [
        ("receipt_number", "String", "Yes", "Printed receipt number (e.g., \"12345\" or \"00123\")"),
        ("count", "Integer", "Yes", "Total number of paid tickets in the receipt (e.g., 1, 2, 3)")
    ]

    for row_idx, data in enumerate(rows_data):
        row = table_fields.rows[row_idx + 1]
        bg = "FFFFFF" if row_idx % 2 == 0 else "F9FAFB"
        for col_idx, text in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if col_idx == 0:
                run.font.name = 'Consolas'
                run.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Server Response
    p_resp_t = doc.add_paragraph()
    p_resp_t.paragraph_format.space_after = Pt(2)
    p_resp_t.add_run("Server Response Format:").bold = True
    
    doc.add_paragraph("Upon successful processing, the backend responds with 200 OK:").paragraph_format.space_after = Pt(3)
    add_code_block('{\n  "status": "ok"\n}')

    p_err = doc.add_paragraph()
    p_err.paragraph_format.space_after = Pt(4)
    p_err.add_run("If the authorization token is missing or invalid, the server responds with ")
    p_err.add_run("401 Unauthorized").bold = True
    p_err.add_run(" and rejects the operation.")

    # Fault tolerance
    p_ft_t = doc.add_paragraph()
    p_ft_t.paragraph_format.space_after = Pt(2)
    p_ft_t.add_run("Fault Tolerance & Retries:").bold = True

    p_idemp = doc.add_paragraph(style='List Bullet')
    p_idemp.paragraph_format.space_after = Pt(2)
    p_idemp.add_run("Idempotency (Duplicate Protection): ").bold = True
    p_idemp.add_run("Transactions are uniquely identified by receipt_number. If a network timeout causes the POS terminal to retry sending the same receipt, the system will not create duplicate activation credits and will safely return 200 OK.")

    p_off = doc.add_paragraph(style='List Bullet')
    p_off.paragraph_format.space_after = Pt(6)
    p_off.add_run("Offline Queue on POS: ").bold = True
    p_off.add_run("If the backend is temporarily unreachable when closing a receipt, the POS module must store the payload in a local queue and automatically retry sending once connectivity is restored.")

    # 2.2 Guest Lookup & Authorization at POS
    add_sub_heading("2.2 Guest Lookup & Authorization at POS")
    p_auth_desc = doc.add_paragraph()
    p_auth_desc.paragraph_format.space_after = Pt(3)
    p_auth_desc.add_run("If the POS terminal needs to retrieve guest data, it sends a GET request to our cloud server:")

    add_code_block("GET https://<HP_SERVER>/api/user")

    p_param_desc = doc.add_paragraph()
    p_param_desc.paragraph_format.space_after = Pt(2)
    p_param_desc.add_run("The request parameter can be either:")

    p_p1 = doc.add_paragraph(style='List Bullet')
    p_p1.paragraph_format.space_after = Pt(1)
    p_p1.add_run("User ID (scanned from the QR code shown to the guest upon registration)")

    p_p2 = doc.add_paragraph(style='List Bullet')
    p_p2.paragraph_format.space_after = Pt(3)
    p_p2.add_run("Phone number")

    p_resp_desc = doc.add_paragraph()
    p_resp_desc.paragraph_format.space_after = Pt(6)
    p_resp_desc.add_run("In response, our cloud returns the guest profile and linked children data.")

    # 2.3 Turnstiles & Access Control System (ACS)
    add_sub_heading("2.3 Turnstiles and Access Control System (ACS / Time Tracking)")
    p_turn_desc = doc.add_paragraph()
    p_turn_desc.paragraph_format.space_after = Pt(3)
    p_turn_desc.add_run("After wristband assignment, our system can send a notification request to a third-party Access Control System (ACS / Turnstiles):")

    p_post_params = doc.add_paragraph()
    p_post_params.paragraph_format.space_after = Pt(2)
    p_post_params.add_run("POST request with parameters:").bold = True

    p_t1 = doc.add_paragraph(style='List Bullet')
    p_t1.paragraph_format.space_after = Pt(1)
    p_t1.add_run("user_id — User / Child identifier")

    p_t2 = doc.add_paragraph(style='List Bullet')
    p_t2.paragraph_format.space_after = Pt(1)
    p_t2.add_run("rfid — RFID wristband number")

    p_t3 = doc.add_paragraph(style='List Bullet')
    p_t3.paragraph_format.space_after = Pt(6)
    p_t3.add_run("date — Registration timestamp / date")

    # 2.4 New Guest Registration Webhook
    add_sub_heading("2.4 New Guest Registration Webhook (CRM Integration)")
    p_crm_desc = doc.add_paragraph()
    p_crm_desc.paragraph_format.space_after = Pt(3)
    p_crm_desc.add_run("Upon registration of a new guest, our server can trigger a webhook on a third-party CRM system to transfer user profile details:")

    p_c1 = doc.add_paragraph(style='List Bullet')
    p_c1.paragraph_format.space_after = Pt(1)
    p_c1.add_run("ID in our system (user_id)")

    p_c2 = doc.add_paragraph(style='List Bullet')
    p_c2.paragraph_format.space_after = Pt(1)
    p_c2.add_run("Contact details (phone number)")

    p_c3 = doc.add_paragraph(style='List Bullet')
    p_c3.paragraph_format.space_after = Pt(6)
    p_c3.add_run("Profile data (full name, gender, age, date of birth)")

    doc.save(output_path)
    print(f"English DOCX successfully compiled at: {output_path}")

if __name__ == '__main__':
    create_docx()
