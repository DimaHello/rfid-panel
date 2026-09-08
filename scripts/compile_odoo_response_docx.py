import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    scripts_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(scripts_dir)
    docs_dir = os.path.join(project_dir, 'docs')
    output_path = os.path.join(docs_dir, 'Ответ_команде_Odoo_по_интеграции_RFID.docx')

    doc = docx.Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Base style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(35, 35, 35)

    def add_title(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(20)
        run.font.bold = True
        run.font.color.rgb = RGBColor(17, 24, 39)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(17, 24, 39)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(31, 41, 55)
        return p

    def add_h4(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(55, 65, 81)
        return p

    def add_callout(text_or_runs, prefix="", bg_color="F9FAFB"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, bg_color)
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if prefix:
            r = p.add_run(prefix)
            r.bold = True
            r.font.color.rgb = RGBColor(17, 24, 39)
        if isinstance(text_or_runs, str):
            p.add_run(text_or_runs)
        elif isinstance(text_or_runs, list):
            for item in text_or_runs:
                if isinstance(item, tuple):
                    r = p.add_run(item[0])
                    if len(item) > 1 and item[1]:
                        r.bold = True
                    if len(item) > 2 and item[2]:
                        r.font.color.rgb = item[2]
                else:
                    p.add_run(str(item))
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    def add_code_block(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = tbl.rows[0].cells[0]
        c.width = Inches(6.8)
        set_cell_background(c, "282C34")
        set_cell_margins(c, top=120, bottom=120, left=160, right=160)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(235, 235, 235)
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(4)

    # --- Title ---
    add_title("Ответ команде интеграции Odoo: RFID-панель и Анкета регистрации Hello Park")

    # --- Список вопросов ---
    add_h2("Список вопросов")
    p_q_intro = doc.add_paragraph("Интеграторы прислали следующий список вопросов для формирования BRD (Business Requirements Document):")
    p_q_intro.paragraph_format.space_after = Pt(4)

    q1 = doc.add_paragraph(style='List Number')
    q1.add_run("RFID Band Binding & Unbinding API Documentation:").bold = True
    q1.paragraph_format.space_after = Pt(1)
    for sub in [
        "Эндпоинты API и авторизация",
        "Спецификация запросов и ответов",
        "Обязательные параметры и правила валидации",
        "Примеры JSON payload (запрос / ответ)",
        "Коды ошибок и логика их обработки",
        "Бизнес-правила и зависимости",
        "Примеры кода интеграции (если есть)"
    ]:
        p_sub = doc.add_paragraph(style='List Bullet 2')
        p_sub.paragraph_format.space_after = Pt(1)
        p_sub.add_run(sub)

    q2 = doc.add_paragraph(style='List Number')
    q2.add_run("Existing Hello Parks Guest Registration Form:").bold = True
    q2.paragraph_format.space_before = Pt(4)
    q2.paragraph_format.space_after = Pt(1)
    for sub in [
        "Форма регистрации гостей, используемая в Hello Park",
        "Детализация всех собираемых полей",
        "Обязательные / опциональные поля и валидация",
        "Согласия, правила парка и условия обработки данных",
        "Логика идентификации и сопоставления клиентов",
        "Учетные данные (доступы / логины) к тестовому стенду (UAT) для RFID-панели и анкеты гостя"
    ]:
        p_sub = doc.add_paragraph(style='List Bullet 2')
        p_sub.paragraph_format.space_after = Pt(1)
        p_sub.add_run(sub)

    # --- ЧАСТЬ 1 ---
    add_h2("ЧАСТЬ 1. Пользовательский путь и сценарии (User Journey & Architecture)")

    add_callout(
        [
            ("Интерактивные демонстрационные стенды (демо-режим, без паролей):\n", True),
            ("• Мобильная анкета гостя (SPA): ", True),
            ("https://aihrtest-create.github.io/rfid-panel/guest-registration/\n", False),
            ("• RFID-панель кассира: ", True),
            ("https://aihrtest-create.github.io/rfid-panel/\n\n", False),
            ("(Интерфейсы развернуты в виде работающих веб-прототипов, их можно протестировать в браузере напрямую).", False)
        ],
        prefix="🔗 ",
        bg_color="FFF7ED"
    )

    add_h3("1.1. Общая концепция взаимодействия систем (как это работает сейчас)")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Регистрация гостей и игровая система: ").bold = True
    p.add_run("Регистрация посетителей, профили детей, их прогресс на интерактивных аттракционах парка, очки и виртуальные Аватары ведутся на стороне системы Hello Park.")
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Касса парка (POS): ").bold = True
    p.add_run("Продажа входных билетов и печать чеков осуществляются в кассовой системе (POS) парка.")
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Контроль выдачи браслетов: ").bold = True
    p.add_run("Физический RFID-браслет выдается ребенку только после подтверждения оплаты билета на кассе. Кассовая система при оплате чека передает сигнал (токен активации) в систему Hello Park, разрешая привязку браслета. Без оплаченного входного билета выдать браслет невозможно.")
    p.paragraph_format.space_after = Pt(6)

    add_h3("1.2. Схема пользовательского пути (User Flow)")
    img_flow = os.path.join(docs_dir, 'screenshots', 'user_flow_diagram.png')
    if os.path.exists(img_flow):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        p_img.add_run().add_picture(img_flow, width=Inches(6.6))

    steps = [
        ("Шаг 1 (Гость): ", "При входе в парк гость сканирует QR-код со своего смартфона и заполняет простую мобильную анкету (номер телефона с подтверждением по СМС, ФИО родителя, имена и даты рождения детей)."),
        ("Шаг 2 (Сервер): ", "Данные семьи автоматически сохраняются в глобальной и локальной базе данных Hello Park и моментально отображаются в веб-интерфейсе RFID-панели на компьютере кассира."),
        ("Шаг 3 (Касса парка / POS): ", "Гость подходит к кассе и оплачивает входные билеты. На этом этапе кассир ещё не может привязать браслет ребенку — в системе нет токена активации."),
        ("Шаг 4 (Интеграция): ", "Сразу после фискализации чека касса парка (POS) отправляет простой HTTP POST запрос (вебхук) на бэкенд Hello Park с номером чека и количеством оплаченных билетов."),
        ("Шаг 5 (Кассир): ", "В RFID-панели моментально начисляется количество доступных активаций (слотов). Кассир выбирает ребенка, прикладывает чистый браслет к настольному считывателю — браслет привязан, ребенок идет играть.")
    ]
    for s_title, s_desc in steps:
        p_step = doc.add_paragraph(style='List Number')
        p_step.paragraph_format.space_after = Pt(3)
        p_step.add_run(s_title).bold = True
        p_step.add_run(s_desc)

    add_h3("1.3. Рабочее место кассира (RFID Cashier Panel)")
    p_cash = doc.add_paragraph("Веб-приложение, работающее на компьютере кассира рядом с POS-терминалом. Интерфейс оптимизирован для быстрой работы без лишних кликов:")
    p_cash.paragraph_format.space_after = Pt(4)

    img_main = os.path.join(docs_dir, 'screenshots', 'main_screen_en.png')
    if os.path.exists(img_main):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        p_img.add_run().add_picture(img_main, width=Inches(6.6))

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Строка родителя: ").bold = True
    p.add_run("ФИО, телефон, список привязанных детей.")
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Статусы ребенка:").bold = True
    p.paragraph_format.space_after = Pt(1)

    for st_title, st_desc in [
        ("NONE (Серый): ", "ребенок зарегистрирован в базе, но браслет еще не выдан."),
        ("BRACELET (Оранжевый): ", "браслет привязан на кассе, ребенок находится в игровой зоне."),
        ("AVATAR (Фиолетовый): ", "ребенок создал игрового персонажа на интерактивных проекциях парка. Вся статистика, баллы и заработанные подарки привязаны к его постоянному профилю.")
    ]:
        p_st = doc.add_paragraph(style='List Bullet 2')
        p_st.paragraph_format.space_after = Pt(1)
        p_st.add_run(st_title).bold = True
        p_st.add_run(st_desc)

    add_h3("1.4. Ключевые сценарии работы с гостями")

    add_h4("Сценарий 1: Первичный визит новой семьи")
    s1_steps = [
        "Родитель заполняет онлайн-анкету при входе.",
        "Подходит к кассе, кассир пробивает билеты в кассовой системе (POS).",
        "Касса шлет сигнал оплаты → в панели Hello Park загораются доступные слоты активации.",
        "Кассир прикладывает браслет к считывателю → физический браслет связывается с ребенком → ребенок идет на аттракционы."
    ]
    for step in s1_steps:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(step)

    add_h4("Сценарий 2: Повторный визит гостя")
    p_desc = doc.add_paragraph("Дети регулярно возвращаются в парк повторно (через неделю, месяц или полгода):")
    p_desc.paragraph_format.space_after = Pt(2)
    s2_steps = [
        "Физический браслет с прошлого посещения сдан на выходе, но игровой прогресс (Аватар, уровень, набранные очки) навсегда сохранен в базе данных Hello Park.",
        "При новом визите кассир находит семью в системе (по номеру телефона или имени).",
        "Кассир берет любой новый чистый браслет, нажимает кнопку «Новый браслет» напротив ребенка и сканирует его.",
        "Новый браслет связывается с постоянным профилем ребенка (child_id) → ребенок прикладывает браслет к аттракционам в парке, и система моментально восстанавливает его прежний игровой уровень и баланс очков."
    ]
    for step in s2_steps:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(step)

    add_h4("Сценарий 3: Возврат браслета и повторное использование (Same-Day Reuse)")
    p_desc = doc.add_paragraph("В парках с высоким трафиком браслеты циркулируют непрерывно:")
    p_desc.paragraph_format.space_after = Pt(2)
    s3_steps = [
        "При выходе из парка ребенок сдает браслет на стойку кассы.",
        "Кассир нажимает кнопку «Очистить браслет» в панели и прикладывает его к считывателю — браслет отвязывается за 1 секунду. При этом игровой прогресс ребенка остается в безопасности.",
        "Очищенный браслет сразу же выдается следующему вошедшему гостю.",
        "Ежедневный ночной сброс: Каждую ночь в 03:00 сервер Hello Park автоматически отвязывает абсолютно все физические метки браслетов от детей. Утром весь фонд браслетов парка снова готов к чистой выдаче."
    ]
    for step in s3_steps:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(step)

    # --- ЧАСТЬ 2 ---
    add_h2("ЧАСТЬ 2. Техническая спецификация API (Binding & POS Integration)")

    add_h3("2.1. Архитектурная развилка: кто выполняет привязку браслета?")
    doc.add_paragraph("Перед фиксацией BRD необходимо согласовать с командой Odoo один из двух вариантов:").paragraph_format.space_after = Pt(4)

    add_h4("Вариант А: RFID-панель Hello Park")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Как работает: ").bold = True
    p.add_run("На кассе парка открыта веб-панель кассира Hello Park, подключен настольный USB RFID-считыватель.")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Роль Odoo: ").bold = True
    p.add_run("Odoo не работает напрямую со считывателем браслетов. Odoo только отправляет вебхук об успешной оплате билета (POST /api/v1/pos/events).")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Результат: ").bold = True
    p.add_run("В панели кассира начисляются токены активации, кассир прикладывает браслет в интерфейсе Hello Park.")
    p.paragraph_format.space_after = Pt(6)

    add_h4("Вариант Б: Прямая привязка из Odoo POS")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Как работает: ").bold = True
    p.add_run("Odoo POS напрямую подключает считыватель браслетов к своему интерфейсу.")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Роль Odoo: ").bold = True
    p.add_run("После сканирования браслета Odoo вызывает внешнее API Hello Park для связки child_id + rfid_uid.")
    p.paragraph_format.space_after = Pt(6)

    add_h3("2.2. Спецификация для Варианта А (Вебхук оплаты из Odoo в Hello Park)")
    doc.add_paragraph("Касса отправляет уведомление сразу после фискализации/оплаты чека, если в чеке есть хотя бы один входной билет.").paragraph_format.space_after = Pt(4)

    p_spec = doc.add_paragraph()
    p_spec.add_run("URL: ").bold = True
    p_spec.add_run("POST https://<HP_SERVER_URL>/api/v1/pos/events\n")
    p_spec.add_run("Заголовки:\n").bold = True
    p_spec.paragraph_format.space_after = Pt(2)

    add_code_block("Content-Type: application/json; charset=utf-8\nAuthorization: Bearer <API_SECRET_KEY>")

    doc.add_paragraph("Тело запроса (JSON):").bold = True
    add_code_block('{\n  "receipt_number": "REC-2026-0904-001",\n  "count": 2\n}')

    p_params = doc.add_paragraph()
    p_params.add_run("Параметры:\n").bold = True
    p_p1 = doc.add_paragraph(style='List Bullet')
    p_p1.add_run("receipt_number (String, Обязательный): ").bold = True
    p_p1.add_run("Уникальный номер чека в кассе Odoo. Используется для идемпотентности (защита от дублирования активаций при повторной отправке из-за сбоя сети).")
    p_p1.paragraph_format.space_after = Pt(1)
    p_p2 = doc.add_paragraph(style='List Bullet')
    p_p2.add_run("count (Integer, Обязательный): ").bold = True
    p_p2.add_run("Количество оплаченных детских входных билетов в чеке.")
    p_p2.paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Успешный ответ (200 OK):").bold = True
    add_code_block('{\n  "status": "ok"\n}')

    doc.add_paragraph("Коды ошибок:").bold = True
    p_e1 = doc.add_paragraph(style='List Bullet')
    p_e1.add_run("401 Unauthorized: ").bold = True
    p_e1.add_run("Неверный или отсутствующий ключ авторизации (API_SECRET_KEY).")
    p_e1.paragraph_format.space_after = Pt(1)
    p_e2 = doc.add_paragraph(style='List Bullet')
    p_e2.add_run("400 Bad Request: ").bold = True
    p_e2.add_run("Пропущены обязательные поля или некорректный JSON.")
    p_e2.paragraph_format.space_after = Pt(6)

    add_h3("2.3. Спецификация для Варианта Б (Прямые методы Bind / Unbind из Odoo)")
    doc.add_paragraph("Если Odoo реализует сканирование браслетов на своей стороне:").paragraph_format.space_after = Pt(4)

    add_h4("Метод 1: Привязка браслета (Bind)")
    p = doc.add_paragraph()
    p.add_run("URL: ").bold = True
    p.add_run("POST https://<HP_SERVER_URL>/api/v1/rfid/bind\n")
    p.add_run("Заголовки: ").bold = True
    p.add_run("Authorization: Bearer <API_SECRET_KEY>, Content-Type: application/json")
    p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("Тело запроса:").bold = True
    add_code_block('{\n  "child_id": "CH-10023",\n  "rfid_uid": "04A1B2C3D4",\n  "receipt_number": "REC-2026-0904-001"\n}')

    doc.add_paragraph("Успешный ответ (200 OK):").bold = True
    add_code_block('{\n  "status": "success",\n  "message": "Wristband successfully bound",\n  "data": {\n    "child_id": "CH-10023",\n    "rfid_uid": "04A1B2C3D4",\n    "avatar_status": "ready",\n    "bound_at": "2026-09-04T14:30:00Z"\n  }\n}')

    doc.add_paragraph("Коды ошибок:").bold = True
    for c_err, desc_err in [
        ("400 Bad Request: ", "Не заполнены обязательные параметры."),
        ("404 Not Found: ", "Ребенок с таким child_id не найден в базе."),
        ("409 Conflict: ", "Браслет с таким rfid_uid уже привязан к другому активному гостю (конфликт меток)."),
        ("422 Unprocessable Entity: ", "По данному чеку исчерпано количество оплаченных билетов.")
    ]:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1)
        p.add_run(c_err).bold = True
        p.add_run(desc_err)

    add_h4("Метод 2: Отвязка браслета при выходе / переиспользовании (Unbind)")
    p = doc.add_paragraph()
    p.add_run("URL: ").bold = True
    p.add_run("POST https://<HP_SERVER_URL>/api/v1/rfid/unbind\n")
    p.add_run("Заголовки: ").bold = True
    p.add_run("Authorization: Bearer <API_SECRET_KEY>, Content-Type: application/json")
    p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph("Тело запроса:").bold = True
    add_code_block('{\n  "rfid_uid": "04A1B2C3D4"\n}')

    doc.add_paragraph("Успешный ответ (200 OK):").bold = True
    add_code_block('{\n  "status": "success",\n  "message": "Wristband unbound and released for reuse"\n}')

    add_h3("2.4. Ключевые бизнес-правила и зависимости")
    rules = [
        ("Разделение физического браслета и игрового профиля: ", "Физический браслет — это лишь временный однодневный идентификатор. Аватар, баллы, уровни и заработанные подарки привязаны к профилю ребенка (child_id) в базе данных Hello Park и никогда не удаляются при отвязке браслета."),
        ("Повторный визит (Returning Guests): ", "При новом визите ребенку выдается другой физический браслет, но запрос привязки должен содержать тот же постоянный child_id."),
        ("Защита от дублей (Конфликты): ", "Один браслет нельзя выдать двум детям одновременно. Система вернет 409 Conflict, если браслет числится за кем-то другим прямо сейчас."),
        ("Ежедневный ночной сброс (Daily Wipe): ", "Каждую ночь все привязки физических чипов обнуляются, чтобы утром все браслеты были чистыми.")
    ]
    for r_t, r_d in rules:
        p = doc.add_paragraph(style='List Number')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(r_t).bold = True
        p.add_run(r_d)

    # --- ЧАСТЬ 3 ---
    add_h2("ЧАСТЬ 3. Анкета регистрации гостя (Guest Registration Form)")

    add_h3("3.1. Структура полей анкеты регистрации")
    doc.add_paragraph("Процесс регистрации разбит на последовательные шаги:").paragraph_format.space_after = Pt(3)

    add_h4("Шаг 1: Контактные данные и согласия")
    p1 = doc.add_paragraph(style='List Bullet')
    p1.add_run("Номер телефона (phone): ").bold = True
    p1.add_run("Обязательное поле. Международный формат (маска и количество цифр определяются страной парка). Верификация через одноразовый 4-значный SMS-код (OTP). В демо-версии принимается любой 4-значный код (кроме 0000).")
    p1.paragraph_format.space_after = Pt(2)

    p2 = doc.add_paragraph(style='List Bullet')
    p2.add_run("Юридическое согласие (Чекбокс): ").bold = True
    p2.add_run("Обязательный чекбокс (без него кнопка «Отправить СМС» заблокирована). Содержит согласие с правилами парка, программой лояльности и политикой обработки персональных данных.")
    p2.paragraph_format.space_after = Pt(4)

    add_h4("Шаг 2: Данные семьи и детей")
    f_items = [
        ("ФИО родителя (fio): ", "Обязательное поле, текстовая строка (Фамилия Имя Отчество)."),
        ("Список детей (children): ", "Обязательно добавить минимум одного ребенка. Есть возможность добавления неограниченного числа детей кнопкой «+ Добавить ребёнка»."),
        ("• Имя ребёнка (name): ", "Обязательное, текстовая строка."),
        ("• Дата рождения (dob): ", "Обязательное, формат ДД.ММ.ГГГГ (или выбор через календарь YYYY-MM-DD). Требуется для подбора подходящих по возрасту игр и квестов."),
        ("Дополнительные поля (настраиваются опционально под парк): ", "«Откуда узнали о парке» (источник привлечения / рекламный канал) и «Город» (город проживания гостя).")
    ]
    for ft, fd in f_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(ft).bold = True
        p.add_run(fd)

    add_h4("Шаг 3: Сохранение данных и отображение в системе (Завершение регистрации)")
    doc.add_paragraph("После отправки заполненной анкеты:").paragraph_format.space_after = Pt(2)
    for res_item in [
        "Данные семьи автоматически сохраняются в глобальной и локальной базе данных парка.",
        "Профиль родителя и детей мгновенно появляется в списке зарегистрированных гостей в веб-интерфейсе RFID-панели кассира.",
        "Кассир находит семью в панели по номеру телефона или ФИО для выдачи и привязки браслетов."
    ]:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(res_item)

    add_callout(
        "Для ряда локаций на финальном шаге анкеты гостю дополнительно генерируется персональный QR-код карты лояльности. В этом случае кассиру не требуется вводить номер телефона или искать семью вручную — он просто сканирует QR-код с экрана телефона гостя 2D-сканером на кассе, и нужная карточка открывается моментально.",
        prefix="Опциональный шаг для некоторых парков (генерация QR-кода): ",
        bg_color="F3F4F6"
    )

    add_h3("3.2. Логика идентификации и сопоставления клиентов (Mapping Logic)")
    m_items = [
        ("Идентификатор родителя (Primary Key): ", "Номер мобильного телефона (phone)."),
        ("Идентификатор ребенка: ", "Постоянный child_id, генерируемый при создании карточки ребенка."),
        ("Связь с Odoo: ", "В Odoo карточка клиента должна содержать внешний идентификатор ребенка либо единый номер телефона семьи, чтобы при повторной продаже билета касса могла однозначно передать правильный child_id.")
    ]
    for mt, md in m_items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(mt).bold = True
        p.add_run(md)

    doc.save(output_path)
    print(f"Successfully created: {output_path}")

if __name__ == '__main__':
    create_document()
