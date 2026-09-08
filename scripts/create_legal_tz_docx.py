from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "ТЗ_юристу_внедрение_RFID_панели_во_все_парки.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "202124"
MUTED = "5F6368"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_YELLOW = "FFF4CE"
PALE_RED = "FCE8E6"
WHITE = "FFFFFF"
BORDER = "C7CDD4"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_DXA, (widths, sum(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Calibri", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr()
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 14, 8),
        "Heading 2": (13, BLUE, 11, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True


def add_numbering_definitions(doc):
    numbering = doc.part.numbering_part.element
    existing_abstract = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    next_abstract = max(existing_abstract, default=-1) + 1
    next_num = max(existing_num, default=0) + 1

    def make(kind, abstract_id, num_id):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        lvl.append(lvl_text)
        lvl_jc = OxmlElement("w:lvlJc")
        lvl_jc.set(qn("w:val"), "left")
        lvl.append(lvl_jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        if kind == "bullet":
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "Symbol")
            fonts.set(qn("w:hAnsi"), "Symbol")
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_num_id = OxmlElement("w:abstractNumId")
        abstract_num_id.set(qn("w:val"), str(abstract_id))
        num.append(abstract_num_id)
        numbering.append(num)

    make("bullet", next_abstract, next_num)
    make("decimal", next_abstract + 1, next_num + 1)
    return next_num, next_num + 1


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)


def add_list_item(doc, text, num_id, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=fill, size="2")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(label + " ")
    set_run_font(r, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if isinstance(value, tuple):
                label, url = value
                add_hyperlink(p, label, url)
            else:
                r = p.add_run(str(value))
                set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Стр. ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run2 = paragraph.add_run()
    run2._r.append(fld_char1)
    run2._r.append(instr_text)
    run2._r.append(fld_char2)
    set_run_font(run2, size=9, color=MUTED)


def add_meta_line(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(label + ": ")
    set_run_font(r, size=10.5, bold=True, color=INK)
    r = p.add_run(value)
    set_run_font(r, size=10.5, color=INK)


def page_break(doc):
    doc.add_page_break()


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    bullet_num, decimal_num = add_numbering_definitions(doc)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("HELLO PARK  |  ПРАВОВАЯ ПОДГОТОВКА RFID-ПАНЕЛИ")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    # First page / memo masthead
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ТЕХНИЧЕСКОЕ ЗАДАНИЕ ЮРИСТУ")
    set_run_font(r, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Правовая подготовка внедрения RFID-панели Hello Park во все парки")
    set_run_font(r, size=24, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("Регистрация семьи, данные детей, RFID-браслеты, игровой профиль, программа лояльности и обмен данными между участниками системы")
    set_run_font(r, size=12.5, color=MUTED)

    add_meta_line(doc, "Заказчик", "Hello Park")
    add_meta_line(doc, "Адресат", "Юрист / юридическая команда, отвечающая за персональные данные и договорную модель сети")
    add_meta_line(doc, "Версия", "Черновик 1.0 от 18.08.2026")
    add_meta_line(doc, "Приоритет", "Подготовить российские парки к запуску; для иностранных парков сформировать отдельный маршрут проверки по местному праву")

    add_callout(
        doc,
        "Главный ожидаемый результат.",
        "До запуска у каждого парка должна быть определена юридическая роль, оформлены необходимые уведомления и договоры, опубликован корректный комплект документов, а в интерфейсах должны быть внедрены согласованные тексты, отдельные согласия и доказуемая фиксация волеизъявления гостя.",
    )

    doc.add_heading("1. Задача и границы работы", level=1)
    p = doc.add_paragraph(
        "Нужно провести правовую подготовку к внедрению единой RFID-системы в парках Hello Park. Система связана не только с выдачей браслета: она создает профиль семьи, обрабатывает данные взрослого и ребенка, сохраняет историю посещений и игровой прогресс, участвует в программе лояльности и может обмениваться данными с кассой, СКУД, CRM, SMS-провайдером и другими сервисами."
    )
    p = doc.add_paragraph(
        "Настоящее ТЗ описывает продукт обычным языком и задает перечень юридических решений и документов. Оно не фиксирует заранее, что каждый парк обязательно является самостоятельным оператором: юридическую роль каждого участника должен определить юрист по фактическим целям, полномочиям, договорам и доступу к данным."
    )
    add_callout(
        doc,
        "Важно.",
        "Термин «сборщик персональных данных» в итоговых документах не использовать. Для российского контура необходимо определить: оператор персональных данных, лицо, обрабатывающее данные по поручению оператора, либо самостоятельные операторы по разным целям.",
        fill=PALE_YELLOW,
    )

    doc.add_heading("2. Как работает система", level=1)
    flow_steps = [
        "Гость в парке сканирует QR-код и открывает мобильную форму регистрации, настроенную для конкретного парка и его юридического лица.",
        "Взрослый вводит номер телефона, получает SMS-код и подтверждает номер. До продолжения система показывает обязательные и необязательные юридические действия.",
        "Взрослый указывает свое ФИО и добавляет одного или нескольких детей: имя и дату рождения. Данные ребенка вводит взрослый, поэтому юрист должен определить форму подтверждения полномочий законного представителя.",
        "Система создает семейный профиль и QR-код / идентификатор программы лояльности. Данные сохраняются в базе и становятся доступны в RFID-панели уполномоченному сотруднику парка.",
        "После оплаты билета касса передает в систему номер чека и количество оплаченных входов либо иной токен активации. Без подтверждения оплаты привязка браслета должна быть заблокирована.",
        "Кассир находит семью, выбирает ребенка, считывает UID RFID-браслета и привязывает браслет к профилю ребенка.",
        "Во время посещения система связывает браслет с игровым профилем: аватаром, прогрессом, достижениями, уровнем, баллами и доступными или выданными призами. Фиксируются дата, время и парк посещения.",
        "При повторном посещении семья находится по ФИО или номеру телефона. Новый браслет может быть связан с прежним игровым профилем, чтобы сохранить прогресс.",
        "Дополнительно планируется или допускается обмен с кассой, СКУД, CRM / системой лояльности, SMS-провайдером, сервисами аналитики и Apple / Google Wallet. Каждый фактический обмен должен быть описан в реестре передач.",
    ]
    for item in flow_steps:
        add_list_item(doc, item, decimal_num)

    doc.add_heading("2.1. Два канала регистрации", level=2)
    add_list_item(doc, "Саморегистрация: взрослый заполняет форму на своем телефоне по QR-коду.", bullet_num)
    add_list_item(doc, "Регистрация кассиром: сотрудник вводит те же данные в панели, если гость не может использовать телефон.", bullet_num)
    add_callout(
        doc,
        "Требование к юристу.",
        "Для регистрации кассиром подготовить равноценный сценарий получения и фиксации согласий. Нельзя считать, что кассир может принять согласие за гостя или пропустить его из-за ручного ввода.",
        fill=PALE_YELLOW,
    )

    doc.add_heading("2.2. Ссылки на продуктовые материалы", level=2)
    p = doc.add_paragraph()
    p.add_run("Пример формы регистрации: ")
    r = p.add_run("[ВСТАВИТЬ ПУБЛИЧНУЮ ТЕСТОВУЮ ССЫЛКУ ПЕРЕД ОТПРАВКОЙ ЮРИСТУ]")
    set_run_font(r, bold=True, color="9B1C1C")
    p = doc.add_paragraph()
    p.add_run("Локальный прототип для продуктовой команды: ")
    r = p.add_run("guest-registration/index.html")
    set_run_font(r, italic=True)
    p = doc.add_paragraph()
    p.add_run("Текущие правила и политики всех парков: ")
    add_hyperlink(p, "https://hello-park.ru/r", "https://hello-park.ru/r")
    p = doc.add_paragraph()
    p.add_run("Техническое описание текущего пользовательского пути: ")
    r = p.add_run("docs/ТЗ_Интеграция_Кассы_Активация_Браслетов.md")
    set_run_font(r, italic=True)

    doc.add_heading("3. Реестр данных, который необходимо подтвердить", level=1)
    p = doc.add_paragraph(
        "Юрист должен проверить каждую строку, определить правовое основание, сроки хранения, состав получателей и допустимость обработки данных ребенка. Если поле не нужно для заявленной цели, его следует исключить из продукта."
    )
    data_rows = [
        ("Взрослый / родитель", "ФИО; номер телефона; статус и время SMS-подтверждения", "Создание и поиск семейного профиля; связь с ребенком; обслуживание в парке", "Основание; обязательность ФИО; срок хранения; порядок идентификации при запросах"),
        ("Ребенок", "Имя; дата рождения; связь со взрослым", "Создание игрового профиля; возрастные сценарии; сохранение прогресса", "Кто дает согласие; подтверждение полномочий представителя; минимизация даты рождения"),
        ("Системные идентификаторы", "ID взрослого, ребенка, семьи; QR / ID карты лояльности", "Связь записей и поиск профиля", "Срок; защита QR от извлечения открытых персональных данных"),
        ("RFID", "UID браслета; статус; история привязки, замены и очистки", "Идентификация игрового профиля во время посещения", "Считать ли UID псевдонимным идентификатором; правила повторного использования браслета"),
        ("Посещения", "Парк; дата и время; привязка браслета; история повторных визитов", "Оказание услуги; восстановление профиля; аналитика", "Разделить операционные сроки и аналитику; определить момент обезличивания / удаления"),
        ("Игровой профиль", "Аватар; уровень; прогресс; достижения; баллы; призы; статус выдачи", "Игровая механика и программа лояльности", "Отдельно определить цели игры и маркетинга; правила удаления без потери обязательного учета"),
        ("Оплата / касса", "Номер чека; количество билетов; токен активации; возврат; время операции", "Разрешение выдачи браслета и защита от дублей", "Определить, какие данные относятся к конкретному гостю и какие сроки задают учетные требования"),
        ("Согласия", "Текст и версия документа; выбранные флажки; дата и время; канал; ID профиля; технические доказательства", "Подтверждение законности обработки и рассылок", "Утвердить состав доказательств и срок их хранения; отзыв должен быть связан с версией согласия"),
        ("Технические данные", "IP, устройство, браузер, журналы входов и действий, ошибки, идентификаторы сессий", "Безопасность, расследование ошибок и инцидентов", "Подтвердить фактический состав; сроки; доступ; отражение в политике"),
        ("Сотрудники парка", "Учетная запись кассира / администратора; парк; роль; журнал действий", "Разграничение доступа и аудит операций", "Основание обработки данных работников; хранение логов; информирование работников"),
        ("Потенциальные данные", "Email, пол, фото, изображение лица, геолокация, cookie / аналитика", "Только если реально включено в финальный продукт", "Не включать «на всякий случай»; фото / лицо отдельно проверить на биометрию и распознавание"),
    ]
    add_table(doc, ["Субъект / блок", "Данные", "Зачем нужны", "Что должен решить юрист"], data_rows, [1600, 2460, 2600, 2700], font_size=8.4)

    doc.add_heading("4. Участники и карта передач", level=1)
    participants = [
        ("Юрлицо парка", "Кассиры и администраторы видят и меняют данные; парк оказывает услугу", "Самостоятельный оператор либо лицо по поручению - определить по целям и договору", "Полное наименование, ИНН / ОГРН, адрес, контакт по ПД, номер записи в реестре РКН"),
        ("Центральная компания / владелец платформы", "Разрабатывает и администрирует систему; хранит общую базу; задает механику", "Оператор, совместно определяющее лицо или обработчик по поручению - дать заключение по каждой цели", "Юрлицо, страна, договоры с парками, права доступа, право использовать данные для аналитики / лояльности"),
        ("Хостинг и резервное копирование", "Серверы, базы, резервные копии", "Обработчик / подрядчик", "Провайдер, страна и адрес ЦОД, субподрядчики, сроки резервных копий, удаление"),
        ("SMS-провайдер", "Телефон, код / статус доставки, журналы", "Обработчик либо самостоятельный получатель", "Договор-поручение, страна, состав логов, срок хранения"),
        ("Кассовое ПО", "Чек, билеты, токен; возможно телефон или user_id", "Источник / получатель данных", "Минимизировать payload; определить договор и ответственность за ошибочные передачи"),
        ("СКУД / турникеты", "user_id, UID RFID, дата и время", "Получатель данных для допуска", "Проверить необходимость всех полей, срок и режим доступа"),
        ("CRM / программа лояльности", "ФИО, телефон, профиль, посещения, баллы", "Получатель / обработчик / самостоятельный оператор", "Цели сервиса и маркетинга разделить; определить участников программы"),
        ("Аналитика, теги, cookies", "Онлайн-идентификаторы, IP, события", "Внешний сервис", "Провайдер, страна, согласие на cookies, локализация и трансграничная передача"),
        ("Apple / Google Wallet", "Идентификатор карты и данные пропуска", "Внешний сервис", "Фактический состав и трансграничная передача; не запускать до отдельного решения"),
    ]
    add_table(doc, ["Участник", "Фактическая роль в процессе", "Юридический вопрос", "Данные для проверки"], participants, [1700, 2600, 2300, 2760], font_size=8.3)

    page_break(doc)
    doc.add_heading("5. Ключевые юридические задачи", level=1)

    doc.add_heading("5.1. Определить операторов по каждой цели", level=2)
    for text in [
        "Составить матрицу целей: посещение и допуск, игровой профиль, программа лояльности, маркетинг, аналитика, поддержка и безопасность.",
        "По каждой цели указать, кто определяет состав данных и операции: парк, центральная компания или оба участника независимо.",
        "Если парк действует только по документированному поручению центрального оператора, подготовить поручение обработки с обязательными условиями. Если парк использует данные для собственных целей, квалифицировать его как самостоятельного оператора по этим целям.",
        "Проверить, допускает ли договорная модель сети единые тексты, или документы должны быть индивидуализированы по каждому юрлицу парка.",
    ]:
        add_list_item(doc, text, bullet_num)

    doc.add_heading("5.2. Уведомления Роскомнадзора", level=2)
    for text in [
        "Для каждого российского юрлица проверить наличие и актуальность записи в реестре операторов персональных данных.",
        "Если уведомление не подавалось либо не покрывает новые цели, категории субъектов, состав данных, способы обработки, место базы и подрядчиков - подготовить подачу или изменение сведений до запуска.",
        "Отдельно проверить обязанность уведомления о трансграничной передаче, если к данным получают доступ иностранные компании или данные передаются иностранным сервисам.",
        "Предоставить Hello Park подтверждение подачи, регистрационный номер и копию окончательной формы по каждому парку.",
    ]:
        add_list_item(doc, text, bullet_num)
    p = doc.add_paragraph()
    p.add_run("Официальная форма уведомлений утверждена приказом Роскомнадзора от 28.10.2022 № 180: ")
    add_hyperlink(p, "официальная публикация", "https://publication.pravo.gov.ru/Document/View/0001202212150022")

    doc.add_heading("5.3. Правовые основания и согласия", level=2)
    for text in [
        "Разнести основания обработки: исполнение договора / оказание услуги, согласие, обязанности по закону и иные применимые основания. Не использовать согласие там, где юрист рекомендует более устойчивое основание.",
        "Подготовить отдельное согласие на обработку персональных данных. С 01.09.2025 согласие должно оформляться отдельно от иной информации и документов, которые подтверждает или подписывает субъект; текущий объединенный флажок формы требует пересмотра.",
        "Отделить обязательную обработку для услуги от необязательных маркетинговых коммуникаций. Отказ от рекламы не должен блокировать вход в парк, создание профиля или начисление игровых баллов, если маркетинг не нужен для услуги.",
        "Подготовить отдельное подтверждаемое согласие на рекламу по SMS, телефону, email и мессенджерам, если такие каналы используются.",
        "Утвердить текст о данных ребенка и подтверждение, что форму заполняет родитель / законный представитель. Определить достаточный способ проверки полномочий без избыточного сбора паспортных данных.",
        "Если профиль, имя, фото, аватар или достижения публикуются для неопределенного круга лиц, отдельно проверить режим данных, разрешенных для распространения.",
    ]:
        add_list_item(doc, text, bullet_num)

    doc.add_heading("5.4. Публичные и договорные документы", level=2)
    for text in [
        "Политика обработки персональных данных для RFID-системы и формы регистрации - по каждому оператору либо единый документ с ясным распределением ролей.",
        "Согласие взрослого на обработку собственных данных.",
        "Согласие / подтверждение законного представителя на обработку данных ребенка.",
        "Правила программы лояльности с описанием владельца программы, участников, баллов, призов, срока действия и последствий удаления профиля.",
        "Отдельное согласие на рекламные сообщения.",
        "Cookie-уведомление и настройка аналитики для веб-формы.",
        "Поручения на обработку и договорные приложения с центральной компанией, парками, хостингом, SMS, CRM, СКУД и другими подрядчиками.",
        "Внутренние документы оператора: политика, перечни данных и допущенных лиц, назначение ответственного, порядок запросов субъектов, сроки хранения и уничтожения, обучение персонала, порядок инцидентов.",
    ]:
        add_list_item(doc, text, bullet_num)

    doc.add_heading("5.5. Защита, доступ, сроки и инциденты", level=2)
    for text in [
        "Утвердить принцип доступа только к данным своего парка и необходимым функциям; отдельно описать доступ центральной поддержки и администраторов.",
        "Задать сроки хранения по каждой цели, включая архивы и резервные копии. Формулировки «бессрочно» или «неограниченно» заменить на обоснованные сроки и события удаления.",
        "Установить процесс исправления, выгрузки, отзыва согласия, прекращения рассылок, блокирования и удаления данных, включая данные ребенка и связанные записи RFID / лояльности.",
        "Подготовить регламент инцидентов: кто фиксирует событие, кому парк сообщает внутри сети, кто взаимодействует с Роскомнадзором и субъектами. Процесс должен позволять выполнить применимые уведомления Роскомнадзора в течение 24 часов об инциденте и в течение 72 часов о результатах внутреннего расследования.",
        "Подтвердить первичную запись и хранение данных граждан РФ в базах на территории России и проверить все иностранные сервисы / доступы.",
    ]:
        add_list_item(doc, text, bullet_num)

    doc.add_heading("6. Что уже видно по текущему прототипу и документам", level=1)
    findings = [
        "Один флажок формы объединяет правила парка, программу лояльности и согласие на обработку данных. Юрист должен дать отдельные тексты и указать, какие флажки обязательны, а какие добровольны.",
        "Прототип жестко ведет на документы парка Avenue Sever / «Селигерская». При запуске ссылка и реквизиты должны определяться конкретным парком, а сервер должен фиксировать версию показанного документа.",
        "Текущие политики парков в основном описывают посетителей сайта и типовой набор «имя, email, телефон, cookies». RFID-система добавляет данные ребенка, UID браслета, посещения, игровой прогресс, баллы, призы и обмены с кассой / СКУД / CRM; действующий охват нужно расширить.",
        "В опубликованных политиках уже указаны разные операторы по паркам. Это подтверждает, что нельзя выпустить одну обезличенную форму без маршрутизации по юрлицу и без общей договорной модели.",
        "Некоторые политики заявляют отсутствие передачи третьим лицам или неограниченный срок обработки. Эти положения могут не соответствовать будущей архитектуре с центральной платформой и подрядчиками и должны быть проверены.",
        "В текущем прототипе регистрация и SMS являются имитацией и данные не сохраняются на сервере. До юридического заключения нужно предоставить финальную архитектуру: где база, кто администрирует, кто видит данные, какие API и подрядчики реально включены.",
    ]
    for f in findings:
        add_list_item(doc, f, bullet_num)
    add_callout(
        doc,
        "Стоп-фактор запуска.",
        "Не переносить текущий единый флажок и ссылки «Селигерской» в промышленную форму для всех парков без письменного согласования юриста.",
        fill=PALE_RED,
    )

    doc.add_heading("7. Результаты, которые должен передать юрист", level=1)
    deliverables = [
        "Письменное заключение по ролевой модели: оператор(ы), обработчики по поручению, самостоятельные цели и ответственность сторон.",
        "Матрица готовности по каждому парку: юрлицо, страна, оператор, запись в реестре / местная регистрация, необходимые действия, статус и блокеры.",
        "Заполненные и готовые к подаче уведомления Роскомнадзора и уведомления об изменении сведений; при необходимости - документы по трансграничной передаче.",
        "Финальный комплект публичных документов и согласий с переменными для конкретного парка.",
        "Договорная схема и шаблоны поручений / приложений со всеми участниками обработки.",
        "Точные требования к интерфейсу: тексты, отдельные флажки, обязательность, порядок переходов, возрастные формулировки, ссылки, версии и состав журналируемых доказательств.",
        "Реестр данных и передач с утвержденными сроками хранения и правилами удаления.",
        "Внутренний пакет для парков: приказ о назначении ответственного, инструкции кассиру и администратору, матрица доступа, обработка запросов, уничтожение, обучение и инциденты.",
        "Финальный юридический чек-лист допуска к запуску с подписью / подтверждением по каждому парку.",
    ]
    _, deliverable_num = add_numbering_definitions(doc)
    for d in deliverables:
        add_list_item(doc, d, deliverable_num)

    doc.add_heading("7.1. Критерии приемки", level=2)
    criteria = [
        "По каждому парку есть однозначно указанное юрлицо и юридическая роль.",
        "Нет ссылок на документы другого парка или другого оператора.",
        "Все поля формы отражены в документах и имеют цель, основание и срок хранения.",
        "Согласия разделены и могут быть доказаны по конкретной версии текста.",
        "Маркетинг отключаем и не является условием обязательной услуги.",
        "Данные ребенка и полномочия взрослого оформлены согласованным способом.",
        "Все получатели и подрядчики отражены в договорной и публичной документации.",
        "Подтверждены уведомления регулятору и локализация базы; трансграничные потоки либо оформлены, либо отключены.",
        "Назначены ответственные и утвержден регламент запросов / удаления / инцидентов.",
        "Продуктовая команда получила точные тексты и правила поведения интерфейса без двусмысленных рекомендаций.",
    ]
    for c in criteria:
        add_list_item(doc, "☐ " + c, bullet_num)

    page_break(doc)
    doc.add_heading("8. Информация, которую Hello Park должен предоставить юристу", level=1)
    inputs = [
        "Перечень всех парков в фактическом охвате запуска и их юридические реквизиты.",
        "Договоры франшизы / лицензии / оказания IT-услуг между центральной компанией и парками.",
        "Существующие номера записей в реестре операторов и копии ранее поданных уведомлений.",
        "Архитектурную схему: серверы, базы, резервные копии, страны размещения, администраторы и удаленный доступ.",
        "Перечень подрядчиков и договоров: хостинг, SMS, CRM, аналитика, касса, СКУД, Wallet, поддержка.",
        "Финальный список полей, API payload и журналов; отдельно - какие функции уже есть и какие только планируются.",
        "Матрицу ролей сотрудников и перечень операций, доступных кассиру, администратору парка и центральной поддержке.",
        "Желаемые сроки сохранения профиля, прогресса, истории визитов, баллов и неиспользованных призов.",
        "Публичную тестовую ссылку на форму и тестовый доступ к RFID-панели.",
        "Контакты ответственных от продукта, IT, безопасности, маркетинга и операционного блока.",
    ]
    for i in inputs:
        add_list_item(doc, i, bullet_num)

    doc.add_heading("9. Матрица решений, которую просим заполнить", level=1)
    decisions = [
        ("Оператор данных взрослого", "Кто и по каким целям?", "Юрлицо, обоснование, документы"),
        ("Оператор данных ребенка", "Совпадает ли с оператором взрослого?", "Юрлицо, основание, форма представительства"),
        ("Центральная база", "Оператор или обработчик?", "Договорная конструкция и ответственность"),
        ("Уведомление РКН", "Новое / изменение / не требуется", "Основание, дата, номер, подтверждение"),
        ("Трансграничная передача", "Есть / нет", "Страны, получатели, уведомление или запрет"),
        ("Обязательное согласие", "Нужно ли и на какие цели?", "Финальный отдельный текст"),
        ("Маркетинг", "Какие каналы и кто рекламораспространитель?", "Отдельное согласие и отзыв"),
        ("Сроки хранения", "По каждому блоку данных", "Срок / событие удаления / исключение"),
        ("Инциденты", "Кто уведомляет и принимает решения?", "Регламент, контакты, сроки"),
        ("Допуск к запуску", "Разрешено / условно / запрещено", "Условия и оставшиеся блокеры"),
    ]
    add_table(doc, ["Вопрос", "Решение юриста", "Что должно быть выдано"], decisions, [2500, 2960, 3900], font_size=9)

    doc.add_heading("10. Текущие документы российских парков", level=1)
    p = doc.add_paragraph()
    p.add_run("Единая страница-реестр: ")
    add_hyperlink(p, "hello-park.ru/r", "https://hello-park.ru/r")
    p.add_run(". Перечень ниже фиксирует опубликованные ссылки на дату подготовки ТЗ; юрист должен проверить фактическую редакцию и юридическое лицо в каждом документе.")

    park_rows = [
        ("Благовещенск", "ИП Никеенко Михаил Николаевич", ("Политика", "https://hello-park.ru/blg/privacy"), [("Правила", "https://hello-park.ru/blg/visiting")]),
        ("Владикавказ", "ООО «ИТ ТЕХГРУПП»", ("Политика", "https://hello-park.ru/vladikavkaz/privacy"), [("Правила", "https://hello-park.ru/vladikavkaz/visiting"), ("Лояльность", "https://hello-park.ru/vladikavkaz/loyalty-rules")]),
        ("Воронеж", "ООО «Карусель»", ("Политика", "https://hello-park.ru/voronezh/privacy"), [("Правила", "https://hello-park.ru/voronezh/pravila-poseshcheniya"), ("Лояльность", "https://hello-park.ru/voronezh/loyalty-rules"), ("Акция", "https://hello-park.ru/voronezh/usloviya-akcii-skidka-10-na-pokupku-biletov-onlayn")]),
        ("Каспийск", "ИП Абуев Шамиль Арсланбекович", ("Политика", "https://hello-park.ru/kaspiysk/privacy"), [("Правила", "https://hello-park.ru/kaspiysk/visiting"), ("Лояльность", "https://hello-park.ru/kaspiysk/loyalty-rules")]),
        ("Авиапарк", "ИП Никеенко Анастасия Сергеевна", ("Политика", "https://hello-park.ru/aviapark/privacy"), [("Правила", "https://hello-park.ru/aviapark/visiting"), ("Лояльность", "https://hello-park.ru/aviapark/loyalty-rules")]),
        ("МЕГА Теплый Стан", "ООО «Хеллоу Парк Мега»", ("Политика", "https://hello-park.ru/mega-ts/privacy"), [("Правила", "https://hello-park.ru/mega-ts/visiting"), ("Лояльность", "https://hello-park.ru/mega-ts/loyalty-rules"), ("Акции", "https://hello-park.ru/mega-ts/bonus")]),
        ("Ривьера", "ООО «Диджитал Эксхибишенс»", ("Политика", "https://hello-park.ru/riviera/privacy"), [("Правила", "https://hello-park.ru/riviera/visiting"), ("Лояльность", "https://hello-park.ru/riviera/loyalty-rules")]),
        ("Avenue Sever", "ООО «Компакленд»", ("Политика", "https://hello-park.ru/seligerskaya/policy"), [("Правила", "https://hello-park.ru/seligerskaya/visiting"), ("Лояльность", "https://hello-park.ru/seligerskaya/loyalty-rules")]),
        ("Сахалин", "ООО «Гейм Мастер»", ("Политика", "https://hello-park.ru/sakhalin/privacy"), [("Правила", "https://hello-park.ru/sakhalin/visiting"), ("Лояльность", "https://hello-park.ru/sakhalin/loyalty-rules")]),
        ("МореМолл, Сочи", "ИП Бутаев О. А.", ("Политика", "https://hello-park.ru/sochi/privacy"), [("Правила", "https://hello-park.ru/sochi/visiting"), ("Лояльность", "https://hello-park.ru/sochi/loyalty")]),
    ]
    # Custom 4-column hyperlink table because each row has two links.
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Парк", "Оператор в текущей политике", "Персональные данные", "Другие документы"]
    for idx, header_text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header_text)
        set_run_font(r, size=8.5, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for park, operator, privacy, docs_links in park_rows:
        cells = table.add_row().cells
        for c in cells:
            set_cell_border(c)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(park)
        set_run_font(r, size=8.5)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(operator)
        set_run_font(r, size=8.5)
        p = cells[2].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_hyperlink(p, privacy[0], privacy[1])
        p = cells[3].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for link_idx, (label, url) in enumerate(docs_links):
            if link_idx:
                sep = p.add_run(" | ")
                set_run_font(sep, size=8.5, color=MUTED)
            add_hyperlink(p, label, url)
    set_table_geometry(table, [1450, 2990, 2100, 2820])
    doc.add_paragraph()
    add_callout(
        doc,
        "Примечание.",
        "В таблице даны отдельные прямые ссылки на опубликованные правила, лояльность и акции. Перед отправкой пакета следует повторно сверить его с hello-park.ru/r, поскольку набор и адреса документов могут измениться.",
    )

    doc.add_heading("11. Иностранные парки", level=1)
    p = doc.add_paragraph(
        "На странице hello-park.ru/r также перечислены парки в Азербайджане, Казахстане, Таджикистане, Узбекистане, Бразилии, Колумбии, Литве и Омане. Для них нельзя механически применять российский пакет. Юрист должен определить местное юрлицо, применимое право, требования к данным детей, регистрации оператора, локализации, международной передаче, cookies и маркетинговым коммуникациям; при необходимости привлечь местного консультанта."
    )
    foreign_rows = [
        ("Баку", "Азербайджан", ("Политика", "https://hello-park.ru/baku/privacy"), ("Правила", "https://hello-park.ru/baku/visiting")),
        ("Атырау", "Казахстан", ("Политика", "https://hello-park.ru/atyrau/privacy"), ("Правила", "https://hello-park.ru/atyrau/visiting")),
        ("Душанбе", "Таджикистан", ("Политика", "https://hello-park.ru/dushanbe/privacy"), ("Правила", "https://hello-park.ru/dushanbe/visiting")),
        ("Ташкент", "Узбекистан", ("Политика", "https://hello-park.ru/tashkent/privacy"), ("Правила", "https://hello-park.ru/tashkent/visiting")),
        ("São Paulo", "Бразилия", ("Политика", "https://hello-park.com/saopaulo/privacy-policy"), ("Правила", "https://hello-park.com/saopaulo/visiting")),
        ("Plaza Central, Bogotá", "Колумбия", ("Политика", "https://hello-park.com/bogota/privacy-policy"), ("Правила", "https://hello-park.com/bogota/visiting")),
        ("Kaunas", "Литва / ЕС", ("Политика", "https://hello-park.com/kaunas/privacy-policy"), ("Правила", "https://hello-park.com/kaunas/visiting")),
        ("Muscat", "Оман", ("Политика", "https://hello-park.com/oman/privacy"), ("Правила", "https://hello-park.com/oman/visiting")),
    ]
    add_table(doc, ["Парк", "Страна / режим", "Политика", "Правила"], foreign_rows, [2100, 2260, 2300, 2700], font_size=9)
    p = doc.add_paragraph()
    p.add_run("Дополнительный документ Bogotá: ")
    add_hyperlink(p, "условия мероприятий и праздников", "https://hello-park.com/bogota/politicas-y-condiciones-para-eventos-y-fiestas")

    doc.add_heading("12. Нормативные ориентиры для проверки юристом", level=1)
    p = doc.add_paragraph(
        "Список не заменяет юридическое заключение и должен быть перепроверен на дату запуска. Он дан, чтобы не потерять ключевые блоки при подготовке документов."
    )
    legal_refs = [
        ("Федеральный закон № 152-ФЗ «О персональных данных»", "Определение оператора; принципы и основания обработки; поручение обработки; согласие; распространение; трансграничная передача; локализация; организационные и технические меры; инциденты; уведомление оператора", "https://ips.pravo.gov.ru/api/ips/legislation/document?baseid=None&hash=98490812b3409e2a8d78a11ca9010f434ea3d9250a11dbbdb78690cd5551bdd6"),
        ("Приказ Роскомнадзора от 28.10.2022 № 180", "Формы уведомлений об обработке, изменении сведений и прекращении обработки", "https://publication.pravo.gov.ru/Document/View/0001202212150022"),
        ("Федеральный закон № 38-ФЗ «О рекламе», статья 18", "Предварительное согласие на рекламу по сетям электросвязи", "https://www.consultant.ru/document/cons_doc_LAW_58968/f892dec1383709792452f18d36e7043306e2be0a/"),
        ("Страница документов Hello Park", "Текущие политики, правила посещения, лояльность и акции по паркам", "https://hello-park.ru/r"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, header_text in enumerate(["Источник", "Что проверить", "Ссылка"]):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header_text)
        set_run_font(r, size=9, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    for title, focus, url in legal_refs:
        cells = table.add_row().cells
        for c in cells:
            set_cell_border(c)
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        set_run_font(r, size=8.6)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(focus)
        set_run_font(r, size=8.6)
        p = cells[2].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_hyperlink(p, "Открыть", url)
    set_table_geometry(table, [2800, 4660, 1900])

    page_break(doc)
    doc.add_heading("13. Формат финального ответа юриста", level=1)
    p = doc.add_paragraph(
        "Просим не ограничиваться комментариями к текстам. Финальный ответ должен содержать готовые документы, таблицу по каждому парку и однозначные инструкции продуктовой команде. Для каждого вопроса использовать статус: «готово», «нужно действие», «блокирует запуск» или «не применимо», с ответственным и сроком."
    )
    add_callout(
        doc,
        "Финальный критерий.",
        "Парк допускается к RFID-запуску только после закрытия обязательных юридических, договорных, регуляторных и продуктовых пунктов для его юрлица и юрисдикции.",
    )

    doc.core_properties.title = "ТЗ юристу — правовая подготовка внедрения RFID-панели Hello Park"
    doc.core_properties.subject = "Персональные данные, RFID, дети, программа лояльности, документы парков"
    doc.core_properties.author = "Hello Park"
    doc.core_properties.keywords = "RFID, персональные данные, Hello Park, юрист, оператор"
    doc.core_properties.comments = "Черновик для юридической проверки. Не является юридическим заключением."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
